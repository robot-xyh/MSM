#!/usr/bin/env python3
"""Build the Word-only scientific-problem report and its four Chinese figures."""

from __future__ import annotations

import math
import re
import warnings
from pathlib import Path
from zipfile import ZipFile

import matplotlib

matplotlib.use("Agg")

warnings.filterwarnings("ignore", message="Unable to import Axes3D.*")

import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from matplotlib.lines import Line2D
from matplotlib.patches import Circle, Ellipse, FancyArrowPatch, FancyBboxPatch, Polygon, Rectangle
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


REPORT_DIR = Path(__file__).resolve().parents[1]
ASSET_DIR = REPORT_DIR / "assets" / "scientific_problem_search_assignment"
OUTPUT = REPORT_DIR / "有限视场协同搜索与不确定信息动态分配科学问题报告_CN.docx"

FIGURES = (
    ASSET_DIR / "01_limited_fov_3d_scene.png",
    ASSET_DIR / "02_probability_search_dynamic_assignment.png",
    ASSET_DIR / "03_complete_algorithm_flow.png",
    ASSET_DIR / "04_simulation_to_lab_roadmap.png",
)

TITLE = "有限视场协同搜索与不确定信息动态分配科学问题报告"
KEY_TECHNOLOGY = "不确定信息驱动的多机搜索与目标分配联合滚动优化技术"

BODY_FONT = "仿宋"
HEADING_FONT = "黑体"
LATIN_FONT = "Times New Roman"
MATH_FONT = "Cambria Math"

INK = "202832"
MUTED = "5C6670"
BLUE = "275D85"
TEAL = "22756B"
GREEN = "4E846B"
AMBER = "C68A2B"
RED = "A64B45"
LIGHT_BLUE = "E8F0F6"
LIGHT_TEAL = "E6F2EF"
LIGHT_AMBER = "F7EFD9"
LIGHT_RED = "F6E7E5"
LIGHT_GRAY = "F3F5F6"
WHITE = "FFFFFF"

CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]")
SENTENCE_END_RE = re.compile(r"[。！？]")
SENSITIVE_PHRASES = (
    "本报告",
    "问题的实质不是",
    "核心不是",
    "通俗地说",
    "简单来说",
    "换言之",
    "写成",
    "统一表述为",
    "不虚构",
    "写作口径",
    "贯通",
    "收口",
    "赋能",
    "打造",
    "不是",
    "而是",
)


def configure_matplotlib() -> None:
    font_path = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
    font_manager.fontManager.addfont(font_path)
    font_name = font_manager.FontProperties(fname=font_path).get_name()
    plt.rcParams.update(
        {
            "font.family": font_name,
            "font.sans-serif": [font_name],
            "axes.unicode_minus": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )


def save_figure(fig: plt.Figure, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=300, facecolor="white")
    plt.close(fig)


def project_scene_point(north: float, east: float, height: float) -> tuple[float, float]:
    """Project a world point to a stable axonometric 2D drawing."""

    n = north / 2400.0
    e = (east + 950.0) / 1900.0
    h = height / 850.0
    return 0.055 + 0.68 * n + 0.20 * e, 0.145 + 0.16 * n - 0.09 * e + 0.55 * h


def draw_uncertainty_ellipsoid(ax, center, radii, color: str) -> None:
    cx, cy = project_scene_point(*center)
    width = 2 * (0.68 * radii[0] / 2400.0 + 0.20 * radii[1] / 1900.0)
    height = 2 * (0.55 * radii[2] / 850.0 + 0.04 * radii[1] / 1900.0)
    outer = Ellipse((cx, cy), width, height, angle=9, facecolor=color, edgecolor=color, linewidth=1.4, linestyle="--", alpha=0.13)
    ax.add_patch(outer)
    for offset, scale in ((-0.26, 0.78), (0.0, 1.0), (0.26, 0.78)):
        ring = Ellipse(
            (cx, cy + offset * height),
            width * scale,
            height * 0.23,
            angle=9,
            fill=False,
            edgecolor=color,
            linewidth=0.85,
            alpha=0.58,
        )
        ax.add_patch(ring)
    ax.plot([cx, cx], [cy - height / 2, cy + height / 2], color=color, linewidth=0.8, alpha=0.55, linestyle=":")


def draw_fov_pyramid(ax, apex, ground_center, half_width, half_depth, color: str) -> None:
    gx, gy, gz = ground_center
    corners_world = [
        (gx - half_width, gy - half_depth, gz),
        (gx + half_width, gy - half_depth, gz),
        (gx + half_width, gy + half_depth, gz),
        (gx - half_width, gy + half_depth, gz),
    ]
    corners = [project_scene_point(*point) for point in corners_world]
    apex_xy = project_scene_point(*apex)
    polygon = Polygon([apex_xy, corners[0], corners[1], corners[2], corners[3]], closed=True, facecolor=color, edgecolor=color, linewidth=1.0, alpha=0.10)
    ax.add_patch(polygon)
    for corner in corners:
        ax.plot([apex_xy[0], corner[0]], [apex_xy[1], corner[1]], color=color, linewidth=0.7, alpha=0.65)
    center_xy = project_scene_point(*ground_center)
    ax.plot([apex_xy[0], center_xy[0]], [apex_xy[1], center_xy[1]], color=color, linewidth=1.8)


def build_scene_figure() -> None:
    fig, ax = plt.subplots(figsize=(12, 7.2), dpi=300)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    ground_corners = [
        project_scene_point(0, -950, 0),
        project_scene_point(2400, -950, 0),
        project_scene_point(2400, 950, 0),
        project_scene_point(0, 950, 0),
    ]
    ax.add_patch(Polygon(ground_corners, closed=True, facecolor="#EEF0ED", edgecolor="#AEB6B1", linewidth=1.2))
    for n in np.linspace(0, 2400, 9):
        start = project_scene_point(n, -950, 0)
        end = project_scene_point(n, 950, 0)
        ax.plot([start[0], end[0]], [start[1], end[1]], color="#CCD1CD", linewidth=0.65)
    for e in np.linspace(-950, 950, 8):
        start = project_scene_point(0, e, 0)
        end = project_scene_point(2400, e, 0)
        ax.plot([start[0], end[0]], [start[1], end[1]], color="#CCD1CD", linewidth=0.65)

    regions = [
        ((1520, -390, 130), (270, 190, 115), "#C68A2B"),
        ((1810, 240, 150), (340, 235, 135), "#A64B45"),
        ((1180, 390, 115), (220, 170, 95), "#275D85"),
    ]
    for center, radii, color in regions:
        draw_uncertainty_ellipsoid(ax, center, radii, color)

    drones = [
        (620, -610, 520),
        (760, 20, 610),
        (500, 610, 470),
    ]
    aim_points = [
        (1500, -370, 40),
        (1780, 210, 45),
        (1190, 390, 35),
    ]
    colors = ["#275D85", "#22756B", "#6E5E9A"]
    labels = ["搜索无人机1", "搜索无人机2", "搜索无人机3"]
    directions = [(230, 80, 15), (250, -45, -10), (210, -70, 10)]
    label_offsets = [(-0.075, 0.065), (-0.028, 0.075), (-0.050, -0.065)]

    for drone, aim, color, label, direction, label_offset in zip(drones, aim_points, colors, labels, directions, label_offsets):
        drone_xy = project_scene_point(*drone)
        ax.scatter(drone_xy[0], drone_xy[1], marker="^", s=150, color=color, edgecolor="white", linewidth=1.1, zorder=8)
        ground_xy = project_scene_point(drone[0], drone[1], 0)
        ax.plot([drone_xy[0], ground_xy[0]], [drone_xy[1], ground_xy[1]], color=color, linewidth=0.8, linestyle=":", alpha=0.70)
        ax.text(drone_xy[0] + label_offset[0], drone_xy[1] + label_offset[1], label, color=color, fontsize=12, fontweight="bold", zorder=9)
        draw_fov_pyramid(ax, drone, aim, 115, 88, color)
        end = (drone[0] + direction[0], drone[1] + direction[1], drone[2] + direction[2])
        end_xy = project_scene_point(*end)
        ax.add_patch(FancyArrowPatch(drone_xy, end_xy, arrowstyle="-|>", mutation_scale=14, linewidth=2.0, color=color, zorder=8))

    target_xy = project_scene_point(1760, 250, 115)
    ax.scatter(target_xy[0], target_xy[1], marker="*", s=205, color="#A64B45", edgecolor="white", linewidth=0.9, zorder=10)
    ax.text(target_xy[0] + 0.018, target_xy[1] + 0.022, "已发现待确认目标", color="#8D3733", fontsize=12, fontweight="bold")
    cue_xy = project_scene_point(1450, -420, 260)
    ax.text(cue_xy[0] - 0.05, cue_xy[1] + 0.03, "中心给出的粗略位置与概率范围", color="#80591B", fontsize=12, fontweight="bold")

    d0 = project_scene_point(*drones[0])
    d1 = project_scene_point(*drones[1])
    ax.plot([d0[0], d1[0]], [d0[1], d1[1]], linestyle="--", color="#7B838A", linewidth=1.3)
    ax.text(
        (d0[0] + d1[0]) / 2 - 0.020,
        (d0[1] + d1[1]) / 2 + 0.015,
        "安全间隔",
        color="#5C6670",
        fontsize=11,
        bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.82, "pad": 0.8},
    )

    origin = project_scene_point(70, -900, 0)
    north_end = project_scene_point(550, -900, 0)
    east_end = project_scene_point(70, -350, 0)
    high_end = project_scene_point(70, -900, 330)
    flow_arrow(ax, origin, north_end, color="#59636B")
    flow_arrow(ax, origin, east_end, color="#59636B")
    flow_arrow(ax, origin, high_end, color="#59636B")
    ax.text(north_end[0] + 0.01, north_end[1], "北向", fontsize=10.5, color="#59636B")
    ax.text(east_end[0] + 0.01, east_end[1] - 0.018, "东向", fontsize=10.5, color="#59636B")
    ax.text(high_end[0] - 0.01, high_end[1] + 0.018, "高度", fontsize=10.5, color="#59636B")

    ax.set_title("有限视场条件下的多机协同搜索任务场景", fontsize=20, fontweight="bold", color="#202832", pad=12)

    legend_handles = [
        Line2D([0], [0], marker="^", color="none", markerfacecolor="#275D85", markeredgecolor="white", markersize=9, label="搜索平台与飞行方向"),
        Line2D([0], [0], color="#22756B", linewidth=2, label="云台视轴与有限视锥"),
        Line2D([0], [0], color="#C68A2B", linewidth=1.5, linestyle="--", label="目标可能区域"),
        Line2D([0], [0], marker="*", color="none", markerfacecolor="#A64B45", markersize=11, label="发现后的待确认目标"),
    ]
    ax.legend(handles=legend_handles, loc="upper left", bbox_to_anchor=(0.00, 0.91), frameon=True, framealpha=0.95, fontsize=10)
    save_figure(fig, FIGURES[0])


def probability_field(xg: np.ndarray, yg: np.ndarray) -> np.ndarray:
    a = np.exp(-(((xg + 0.65) / 1.08) ** 2 + ((yg - 0.28) / 0.82) ** 2) / 2)
    b = 0.54 * np.exp(-(((xg - 1.35) / 0.75) ** 2 + ((yg + 1.05) / 0.58) ** 2) / 2)
    c = 0.23 * np.exp(-(((xg - 1.85) / 0.45) ** 2 + ((yg - 1.35) / 0.48) ** 2) / 2)
    p = a + b + c
    return p / p.sum()


def draw_probability_panel(ax, field, title, footprint, target=None, note=None, vmax=None):
    image_artist = ax.imshow(field, origin="lower", extent=(-3, 3, -3, 3), cmap="YlOrRd", vmin=0, vmax=vmax, interpolation="bilinear")
    ax.contour(np.linspace(-3, 3, field.shape[1]), np.linspace(-3, 3, field.shape[0]), field, levels=6, colors="#7A3430", linewidths=0.45, alpha=0.48)
    rect = Rectangle((footprint[0], footprint[1]), footprint[2], footprint[3], fill=False, edgecolor="#22756B", linewidth=2.1, linestyle="--")
    ax.add_patch(rect)
    ax.text(footprint[0] + 0.06, footprint[1] + 0.24, "本轮实际视场", color="#155D55", fontsize=11, fontweight="bold")
    if target is not None:
        ax.scatter(target[0], target[1], marker="*", s=170, color="#275D85", edgecolor="white", linewidth=0.9, zorder=5)
        ax.annotate("连续发现", xy=target, xytext=(target[0] + 0.35, target[1] + 1.05), fontsize=11, fontweight="bold", color="#275D85", arrowprops={"arrowstyle": "->", "color": "#275D85", "lw": 1.5})
    if note:
        ax.text(0.03, 0.04, note, transform=ax.transAxes, fontsize=10.5, color="#202832", bbox={"boxstyle": "round,pad=0.28", "facecolor": "white", "edgecolor": "#D8DDE0", "alpha": 0.94})
    ax.set_title(title, fontsize=16, fontweight="bold", color="#202832", pad=8)
    ax.set_xticks([])
    ax.set_yticks([])
    for spine in ax.spines.values():
        spine.set_color("#B9C0C5")
        spine.set_linewidth(0.8)
    return image_artist


def build_probability_assignment_figure() -> None:
    fig, axes = plt.subplots(2, 2, figsize=(12, 8.0), dpi=300)
    x = np.linspace(-3, 3, 150)
    y = np.linspace(-3, 3, 150)
    xg, yg = np.meshgrid(x, y)
    prior = probability_field(xg, yg)

    footprint = (-1.65, -0.52, 1.80, 1.45)
    mask = (xg >= footprint[0]) & (xg <= footprint[0] + footprint[2]) & (yg >= footprint[1]) & (yg <= footprint[1] + footprint[3])
    no_detection = prior * np.where(mask, 0.16, 1.0)
    no_detection /= no_detection.sum()
    likelihood = 0.02 + np.exp(-(((xg - 1.24) / 0.38) ** 2 + ((yg + 0.96) / 0.32) ** 2) / 2)
    detected = prior * likelihood
    detected /= detected.sum()
    vmax = max(prior.max(), no_detection.max(), detected.max())

    artist = draw_probability_panel(
        axes[0, 0],
        prior,
        "① 初始动态概率网格",
        footprint,
        note="颜色越深，目标存在概率越高",
        vmax=vmax,
    )
    draw_probability_panel(
        axes[0, 1],
        no_detection,
        "② 本轮未发现后的更新",
        footprint,
        note="只降低真实看过区域的概率",
        vmax=vmax,
    )
    draw_probability_panel(
        axes[1, 0],
        detected,
        "③ 连续发现后的更新",
        (0.52, -1.72, 1.72, 1.35),
        target=(1.24, -0.96),
        note="概率向观测方向收敛，其余任务继续滚动",
        vmax=vmax,
    )

    ax = axes[1, 1]
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.set_title("④ 发现后形成稀疏候选关系并一一分配", fontsize=16, fontweight="bold", color="#202832", pad=8)

    drone_y = [0.78, 0.58, 0.38, 0.18]
    target_y = [0.72, 0.47, 0.22]
    for idx, yv in enumerate(drone_y, 1):
        ax.add_patch(Circle((0.18, yv), 0.055, facecolor="#E8F0F6", edgecolor="#275D85", linewidth=1.8))
        ax.text(0.18, yv, f"无人机{idx}", ha="center", va="center", fontsize=10.5, color="#1F4E70", fontweight="bold")
    for idx, yv in enumerate(target_y, 1):
        ax.add_patch(Circle((0.82, yv), 0.055, facecolor="#F6E7E5", edgecolor="#A64B45", linewidth=1.8))
        ax.text(0.82, yv, f"目标{idx}", ha="center", va="center", fontsize=10.5, color="#7F3935", fontweight="bold")

    candidate_edges = [
        (0, 0),
        (0, 1),
        (1, 0),
        (1, 2),
        (2, 1),
        (2, 2),
        (3, 2),
    ]
    chosen_edges = {(0, 1), (1, 0), (3, 2)}
    for di, tj in candidate_edges:
        selected = (di, tj) in chosen_edges
        ax.plot(
            [0.235, 0.765],
            [drone_y[di], target_y[tj]],
            color="#4E846B" if selected else "#AEB6BB",
            linewidth=3.3 if selected else 1.1,
            alpha=1.0 if selected else 0.72,
            zorder=0,
        )

    ax.text(0.04, 0.91, "可用资源", fontsize=12, color="#275D85", fontweight="bold")
    ax.text(0.74, 0.91, "待分目标", fontsize=12, color="#A64B45", fontweight="bold")
    ax.text(
        0.50,
        0.06,
        "规则代价形成基线  →  图网络或强化学习仅作有界修正  →  确定性求解器发布结果",
        ha="center",
        va="center",
        fontsize=10.2,
        color="#202832",
        bbox={"boxstyle": "round,pad=0.34", "facecolor": "#F3F5F6", "edgecolor": "#B9C0C5"},
    )
    ax.legend(
        handles=[
            Line2D([0], [0], color="#AEB6BB", linewidth=1.5, label="通过准入的候选关系"),
            Line2D([0], [0], color="#4E846B", linewidth=3.4, label="最终一一任务关系"),
        ],
        loc="upper center",
        bbox_to_anchor=(0.50, 0.00),
        frameon=False,
        fontsize=9.5,
        ncol=2,
    )

    cbar_ax = fig.add_axes([0.075, 0.073, 0.405, 0.018])
    cbar = fig.colorbar(artist, cax=cbar_ax, orientation="horizontal")
    cbar.set_ticks([0, vmax])
    cbar.set_ticklabels(["低", "高"])
    cbar.set_label("目标存在概率", fontsize=10, labelpad=1)
    cbar.ax.tick_params(labelsize=8)
    fig.suptitle("概率搜索与发现后动态分配原理", fontsize=21, fontweight="bold", color="#202832", y=0.985)
    fig.subplots_adjust(left=0.055, right=0.975, top=0.91, bottom=0.14, wspace=0.16, hspace=0.24)
    save_figure(fig, FIGURES[1])


def flow_box(ax, x, y, w, h, text, face, edge, fontsize=13, linewidth=1.5, text_color="#202832"):
    face = f"#{face}" if not face.startswith("#") else face
    edge = f"#{edge}" if not edge.startswith("#") else edge
    patch = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle="round,pad=0.008,rounding_size=0.012",
        facecolor=face,
        edgecolor=edge,
        linewidth=linewidth,
    )
    ax.add_patch(patch)
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fontsize, color=text_color, linespacing=1.35, fontweight="bold" if "输入" in text or "确定性" in text else "normal")
    return patch


def flow_arrow(ax, start, end, color="#59636B", connectionstyle="arc3", label=None, label_offset=(0, 0)):
    arrow = FancyArrowPatch(
        start,
        end,
        arrowstyle="-|>",
        mutation_scale=14,
        linewidth=1.55,
        color=color,
        connectionstyle=connectionstyle,
        shrinkA=3,
        shrinkB=3,
    )
    ax.add_patch(arrow)
    if label:
        mx = (start[0] + end[0]) / 2 + label_offset[0]
        my = (start[1] + end[1]) / 2 + label_offset[1]
        ax.text(mx, my, label, fontsize=10.5, color=color, fontweight="bold", ha="center", va="center", bbox={"facecolor": "white", "edgecolor": "none", "pad": 0.8})


def build_flow_figure() -> None:
    fig, ax = plt.subplots(figsize=(12, 8.4), dpi=300)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    ax.text(0.5, 0.966, "不确定信息驱动的搜索与目标分配完整算法流程", ha="center", va="center", fontsize=21, fontweight="bold", color="#202832")
    headers = [(0.035, "态势与概率", BLUE), (0.365, "滚动搜索", TEAL), (0.695, "发现后动态分配", RED)]
    for x, label, color in headers:
        ax.add_patch(Rectangle((x, 0.895), 0.27, 0.042, facecolor=f"#{color}", edgecolor="none"))
        ax.text(x + 0.135, 0.916, label, ha="center", va="center", fontsize=14, color="white", fontweight="bold")

    left_boxes = [
        (0.045, 0.775, 0.25, 0.085, "输入：粗略位置、速度、\n不确定范围与资源状态", LIGHT_BLUE, BLUE),
        (0.045, 0.635, 0.25, 0.085, "按统一时刻外推\n建立动态概率网格", LIGHT_BLUE, BLUE),
        (0.045, 0.495, 0.25, 0.085, "生成候选观察单元\n与平台—云台组合", LIGHT_BLUE, BLUE),
        (0.045, 0.245, 0.25, 0.092, "未发现：仅降低\n真实视锥覆盖区域概率", LIGHT_AMBER, AMBER),
    ]
    for args in left_boxes:
        flow_box(ax, *args, fontsize=13)

    middle_boxes = [
        (0.375, 0.775, 0.25, 0.085, "计算搜索收益\n信息、时间、云台、重复、冲突", LIGHT_TEAL, TEAL),
        (0.375, 0.635, 0.25, 0.085, "规则基线形成搜索顺序\n学习方法仅有界调优先级", LIGHT_TEAL, TEAL),
        (0.375, 0.495, 0.25, 0.085, "确定性求解飞行方向\n与云台指向候选", LIGHT_TEAL, TEAL),
        (0.375, 0.355, 0.25, 0.085, "安全检查后执行观察\n记录真实视锥与时间", LIGHT_GRAY, "6B747B"),
    ]
    for args in middle_boxes:
        flow_box(ax, *args, fontsize=12.7)

    diamond = Polygon([(0.50, 0.31), (0.592, 0.235), (0.50, 0.16), (0.408, 0.235)], closed=True, facecolor="#FFF8E6", edgecolor="#C68A2B", linewidth=1.7)
    ax.add_patch(diamond)
    ax.text(0.50, 0.235, "是否形成\n连续发现？", ha="center", va="center", fontsize=13, color="#6E4B16", fontweight="bold")

    right_boxes = [
        (0.705, 0.775, 0.25, 0.085, "建立局部航迹并连续确认\n形成待分目标集合", LIGHT_RED, RED),
        (0.705, 0.635, 0.25, 0.085, "构建稀疏无人机—目标图\n条件筛选并计算规则代价", LIGHT_RED, RED),
        (0.705, 0.495, 0.25, 0.085, "图网络或强化学习\n只作有界候选代价修正", LIGHT_RED, RED),
        (0.705, 0.355, 0.25, 0.085, "匈牙利或最小费用流\n确定最终一一任务关系", LIGHT_RED, RED),
        (0.705, 0.205, 0.25, 0.095, "发布版本与有效期\n任务保持、事件触发重算", LIGHT_GRAY, "6B747B"),
    ]
    for args in right_boxes:
        flow_box(ax, *args, fontsize=12.5)

    flow_arrow(ax, (0.17, 0.775), (0.17, 0.72))
    flow_arrow(ax, (0.17, 0.635), (0.17, 0.58))
    flow_arrow(ax, (0.295, 0.537), (0.375, 0.817))
    flow_arrow(ax, (0.50, 0.775), (0.50, 0.72))
    flow_arrow(ax, (0.50, 0.635), (0.50, 0.58))
    flow_arrow(ax, (0.50, 0.495), (0.50, 0.44))
    flow_arrow(ax, (0.50, 0.355), (0.50, 0.31))
    flow_arrow(ax, (0.408, 0.235), (0.295, 0.291), color="#C68A2B", label="未发现", label_offset=(0, 0.025))
    flow_arrow(ax, (0.17, 0.337), (0.17, 0.635), color="#C68A2B", connectionstyle="arc3,rad=-0.10")
    flow_arrow(ax, (0.592, 0.235), (0.705, 0.817), color="#A64B45", connectionstyle="arc3,rad=-0.18", label="发现", label_offset=(0.015, 0.02))
    flow_arrow(ax, (0.83, 0.775), (0.83, 0.72))
    flow_arrow(ax, (0.83, 0.635), (0.83, 0.58))
    flow_arrow(ax, (0.83, 0.495), (0.83, 0.44))
    flow_arrow(ax, (0.83, 0.355), (0.83, 0.30))
    flow_arrow(ax, (0.705, 0.252), (0.625, 0.677), color="#6B747B", connectionstyle="arc3,rad=0.25")

    boundary = FancyBboxPatch((0.045, 0.055), 0.91, 0.075, boxstyle="round,pad=0.012,rounding_size=0.012", facecolor="#F4F4F1", edgecolor="#A64B45", linewidth=1.7)
    ax.add_patch(boundary)
    ax.text(
        0.50,
        0.093,
        "学习方法权限边界：只调整搜索优先级或可行候选代价；不直接控制平台飞行或云台；不发布最终任务关系",
        ha="center",
        va="center",
        fontsize=12.2,
        color="#7F3935",
        fontweight="bold",
    )
    save_figure(fig, FIGURES[2])


def roadmap_stage(ax, x, title, subtitle, body, color, light_color):
    width = 0.215
    ax.add_patch(FancyBboxPatch((x, 0.25), width, 0.58, boxstyle="round,pad=0.008,rounding_size=0.012", facecolor=light_color, edgecolor=color, linewidth=1.6))
    ax.add_patch(Rectangle((x, 0.72), width, 0.11, facecolor=color, edgecolor=color))
    ax.text(x + width / 2, 0.785, title, ha="center", va="center", fontsize=15, color="white", fontweight="bold")
    ax.text(x + width / 2, 0.735, subtitle, ha="center", va="center", fontsize=11.5, color="white", fontweight="bold")
    ax.text(x + 0.018, 0.69, body, ha="left", va="top", fontsize=12.8, color="#202832", linespacing=1.38)


def build_roadmap_figure() -> None:
    fig, ax = plt.subplots(figsize=(12, 6.7), dpi=300)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.5, 0.94, "从质点训练到封闭实验室的分阶段实施路线", ha="center", va="center", fontsize=21, fontweight="bold", color="#202832")

    roadmap_stage(
        ax,
        0.025,
        "阶段一",
        "质点快速训练",
        "核心工作\n1）批量随机场景\n2）规则与学习并行对照\n3）搜索、分配分项验证\n\n量测输出\n发现率、首次发现\n概率校准、时延\n安全事件计数",
        "#275D85",
        "#E8F0F6",
    )
    roadmap_stage(
        ax,
        0.270,
        "阶段二",
        "AirSim验证",
        "核心工作\n1）真实视锥、平台运动\n2）云台到位与取帧\n3）冲突、时延注入\n\n量测输出\n实际覆盖、执行偏差\n闭环日志、重放一致",
        "#22756B",
        "#E6F2EF",
    )
    roadmap_stage(
        ax,
        0.515,
        "阶段三",
        "封闭实验室验证",
        "规模顺序\n1）2对2、4对4、6对6\n2）固定场景配对复验\n3）相机、云台、网络标定\n\n量测输出\n发现、分配、任务保持\n时延、冲突、失效恢复",
        "#C68A2B",
        "#F7EFD9",
    )
    roadmap_stage(
        ax,
        0.760,
        "阶段四",
        "后续扩规模",
        "扩展条件\n1）6对6达标后启动\n2）逐步增加目标、资源\n3）遮挡、故障、通信退化\n\n量测输出\n规模、时延、效果曲线\n泛化结果、适用范围",
        "#A64B45",
        "#F6E7E5",
    )

    for start in [0.240, 0.485, 0.730]:
        flow_arrow(ax, (start, 0.855), (start + 0.030, 0.855), color="#59636B")

    ax.add_patch(FancyBboxPatch((0.025, 0.095), 0.95, 0.09, boxstyle="round,pad=0.010,rounding_size=0.012", facecolor="#F3F5F6", edgecolor="#7B838A", linewidth=1.4))
    ax.text(
        0.50,
        0.140,
        "统一阶段门：安全违规为零、数据可追溯、与规则基线完成同场景配对比较；未达标则回退定位问题，不进入下一阶段",
        ha="center",
        va="center",
        fontsize=12.2,
        color="#202832",
        fontweight="bold",
    )
    ax.text(0.50, 0.045, "试验前固定场景配置、随机种子、参数版本、输入摘要、执行日志和评价口径", ha="center", va="center", fontsize=11.2, color="#5C6670")
    save_figure(fig, FIGURES[3])


def build_figures() -> None:
    configure_matplotlib()
    build_scene_figure()
    build_probability_assignment_figure()
    build_flow_figure()
    build_roadmap_figure()


def set_run_font(run, *, size: float, bold: bool = False, color: str = INK, east_asia: str = BODY_FONT, latin: str = LATIN_FONT) -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), latin)
    rpr.rFonts.set(qn("w:hAnsi"), latin)
    rpr.rFonts.set(qn("w:eastAsia"), east_asia)


def configure_section(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.35)
    section.bottom_margin = Cm(2.15)
    section.left_margin = Cm(2.45)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.9)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(14)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(28)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    for level, size, before, after in ((1, 18, 16, 8), (2, 16, 12, 5)):
        style = document.styles[f"Heading {level}"]
        style.font.name = LATIN_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(30 if level == 1 else 28)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    equation = document.styles.add_style("公式", WD_STYLE_TYPE.PARAGRAPH)
    equation.font.name = MATH_FONT
    equation.font.size = Pt(12.5)
    equation._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.first_line_indent = Cm(0)
    equation.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    equation.paragraph_format.line_spacing = Pt(23)
    equation.paragraph_format.space_before = Pt(5)
    equation.paragraph_format.space_after = Pt(4)
    equation.paragraph_format.keep_together = True

    caption = document.styles.add_style("图题", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = LATIN_FONT
    caption.font.size = Pt(11)
    caption._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    caption.paragraph_format.line_spacing = Pt(20)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.keep_together = True

    table_title = document.styles.add_style("表题", WD_STYLE_TYPE.PARAGRAPH)
    table_title.font.name = LATIN_FONT
    table_title.font.size = Pt(11.5)
    table_title.font.bold = True
    table_title._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), HEADING_FONT)
    table_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_title.paragraph_format.first_line_indent = Cm(0)
    table_title.paragraph_format.space_before = Pt(6)
    table_title.paragraph_format.space_after = Pt(5)
    table_title.paragraph_format.keep_with_next = True


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def reset_header_footer(container):
    paragraphs = list(container.paragraphs)
    first = paragraphs[0]
    clear_paragraph(first)
    for paragraph in paragraphs[1:]:
        container._element.remove(paragraph._element)
    return first


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))


def add_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = reset_header_footer(section.header)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run(TITLE), size=9, color=MUTED)
    footer = reset_header_footer(section.footer)
    add_page_number(footer)
    pg_num_type = OxmlElement("w:pgNumType")
    pg_num_type.set(qn("w:start"), "1")
    section._sectPr.append(pg_num_type)


def add_cover(document: Document) -> None:
    for _ in range(4):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    title.paragraph_format.line_spacing = Pt(42)
    set_run_font(title.add_run("有限视场协同搜索与\n不确定信息动态分配科学问题报告"), size=26, bold=True, east_asia=HEADING_FONT)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Cm(0)
    subtitle.paragraph_format.space_before = Pt(15)
    set_run_font(subtitle.add_run("科学问题、算法路线与分阶段验证方案"), size=15, bold=True, color=BLUE, east_asia=HEADING_FONT)

    for _ in range(7):
        document.add_paragraph()
    owner = document.add_paragraph()
    owner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    owner.paragraph_format.first_line_indent = Cm(0)
    set_run_font(owner.add_run("项目组"), size=14, east_asia=HEADING_FONT)
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.first_line_indent = Cm(0)
    set_run_font(date.add_run("2026年8月"), size=13)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.first_line_indent = Cm(0)
    note.paragraph_format.space_before = Pt(10)
    set_run_font(note.add_run("科研论证与阶段实施材料"), size=11, color=MUTED)


def add_heading(document: Document, text: str, *, level: int, page_break_before: bool = False) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run(text)
    set_run_font(run, size=18 if level == 1 else 16, bold=True, east_asia=HEADING_FONT)


def add_body(document: Document, text: str):
    paragraph = document.add_paragraph(style="Normal")
    set_run_font(paragraph.add_run(text), size=14)
    return paragraph


def add_key_technology_paragraph(document: Document) -> None:
    paragraph = document.add_paragraph(style="Normal")
    set_run_font(paragraph.add_run("针对上述问题，拟重点突破“"), size=14)
    set_run_font(paragraph.add_run(KEY_TECHNOLOGY), size=14, bold=True, color=BLUE, east_asia=HEADING_FONT)
    set_run_font(
        paragraph.add_run(
            "”。该技术以动态概率网格表示目标可能区域，在同一周期联合计算平台到达、云台可达、信息收益、重复覆盖和航路冲突。目标发现后，先建立通过时效、可达、通信和安全条件检查的稀疏候选关系，再经规则代价、受限学习修正和确定性一一求解形成版本化计划。学习方法只修正搜索优先级或候选代价，最终飞行方向、云台指向和一一任务关系由确定性算法与安全检查决定。"
        ),
        size=14,
    )


def add_equation(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="公式")
    set_run_font(paragraph.add_run(text), size=12.5, latin=MATH_FONT)
    ppr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F4F6F7")
    ppr.append(shading)


def add_figure(document: Document, path: Path, caption: str, number: int, width_cm: float = 16.0) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(0)
    # Inline pictures must not inherit the body's fixed 28 pt line height,
    # otherwise LibreOffice clips the drawing down to an empty text line.
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    caption_paragraph = document.add_paragraph(style="图题")
    set_run_font(caption_paragraph.add_run(f"图 {number}  {caption}"), size=11)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table_title(document: Document, number: int, title: str) -> None:
    paragraph = document.add_paragraph(style="表题")
    set_run_font(paragraph.add_run(f"表 {number}  {title}"), size=11.5, bold=True, east_asia=HEADING_FONT)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    header_row = table.rows[0]
    tr_pr = header_row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for index, value in enumerate(headers):
        cell = header_row.cells[index]
        cell.width = Cm(widths_cm[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, BLUE)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(18)
        set_run_font(paragraph.add_run(value), size=10.5, bold=True, color=WHITE, east_asia=HEADING_FONT)

    for row_index, values in enumerate(rows, 1):
        row = table.add_row()
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.width = Cm(widths_cm[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, "F6F8F9" if row_index % 2 == 0 else WHITE)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            paragraph.paragraph_format.line_spacing = Pt(18)
            paragraph.paragraph_format.space_after = Pt(0)
            set_run_font(paragraph.add_run(value), size=10.2)

    spacer = document.add_paragraph()
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    spacer.paragraph_format.line_spacing = Pt(5)


def enable_picture_quality(document: Document) -> None:
    settings = document.settings._element
    if settings.find(qn("w:doNotCompressPictures")) is None:
        settings.append(OxmlElement("w:doNotCompressPictures"))
    if settings.find(qn("w:updateFields")) is None:
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        settings.append(update_fields)


def build_document() -> None:
    document = Document()
    configure_section(document.sections[0])
    configure_styles(document)
    add_cover(document)

    body_section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section)
    add_header_footer(body_section)

    add_heading(document, "一、任务、主要难点与关键技术", level=1)
    add_heading(document, "（一）任务与科学问题", level=2)
    add_body(
        document,
        "中心节点提供目标的大致位置、运动方向、信息时刻和不确定范围，搜索无人机依靠有限视场光电在局部空域获取证据。多架无人机需要在同一周期确定飞行方向、云台指向、区域搜索顺序和发现后的负责关系。上述决策相互制约，分别处理会使搜索时间消耗在无法按时到达、云台无法稳定指向或友机已经覆盖的区域。",
    )
    add_body(
        document,
        "需要解决的科学问题是在信息持续变化、观察范围受限和资源数量有限的条件下，选择下一次最有价值且能够执行的观察。平台运动决定未来可见范围，云台转动决定当前视锥，发现与未发现共同改变各区域的目标存在概率。飞行、观察和任务关系必须在同一闭环内滚动计算，并始终保留可核验的规则依据、稳定机制和安全约束。",
    )
    add_body(
        document,
        "图1表示平台、云台视锥与目标可能区域的三维关系。中心线索具有高度、横向范围和概率分布，多机接近过程中，平台航向与云台视轴按照各自约束变化。发现尚未确认身份的目标后，还需依据到达能力、观察连续性和证据成熟度重新计算任务关系，局部发现只作为后续确认输入。",
    )
    add_figure(document, FIGURES[0], "有限视场条件下的多机协同搜索三维任务场景", 1)

    add_heading(document, "（二）主要难点", level=2)
    add_body(
        document,
        "第一项难点是不确定范围的统一表示与连续更新。粗略线索随时间扩散，量测时刻与消息到达时刻存在差异，目标机动还会改变概率区域。发现信息使概率向观测一致区域集中，未发现信息只降低真实覆盖且具备探测条件的区域概率，防止一次无目标画面错误消除线索。",
    )
    add_body(
        document,
        "第二项难点是平台与云台机动的联合约束。高概率区域可能要求平台转弯、爬升或等待云台稳定，最近的平台未必最早形成有效观察。多机接近相邻区域时还会发生重复搜索、航路交叉和间隔冲突，搜索收益需同时计入信息价值、到达时间、云台余量、重复覆盖和航路风险。",
    )
    add_body(
        document,
        "第三项难点是局部发现到正式任务关系之间的证据确认。有限视场内先形成待确认的局部航迹，其来源可能是中心线索、邻近目标或重复发现。若按最近距离立即分配，容易造成重复负责、观察中断和频繁换令，因此需先构建稀疏候选关系，再经规则代价、受限修正和一一求解形成任务。",
    )
    add_body(
        document,
        "第四项难点是动态重分配的稳定性与安全性。目标新增、资源故障、任务不可达和计划过期需要及时重算，位置抖动、单次通信迟到和代价小幅变化由保持机制吸收。计划包含保持时间、改善门槛、递增版本、有效期和回执，并由独立安全检查核验禁入区、机间间隔、资源唯一性、通信与能源条件。",
    )
    add_body(
        document,
        "第五项难点是学习方法的可信使用与权限控制。学习方法可估计多周期资源取舍，但训练分布变化、传感器误差和设备时延会引起偏差。其权限限于修正搜索优先级或已准入候选代价，平台控制、云台控制、任务发布和安全判定均由确定性链路完成。",
    )

    add_heading(document, "（三）关键技术", level=2)
    add_key_technology_paragraph(document)

    add_heading(document, "二、具体算法步骤与实施流程", level=1, page_break_before=True)
    add_heading(document, "（一）动态概率搜索与滚动决策", level=2)
    add_body(
        document,
        "每个滚动周期在统一决策时刻形成态势快照，保留线索位置、速度、表示误差范围的协方差、量测时刻、消息到达时刻和来源质量，并读取平台、任务、云台与通信状态。随后依据目标运动模型传播上一周期概率，对过旧或异常信息降低权重。概率网格可使用三维体素或分层平面，其尺度由预测误差、实际视场足迹和计算预算确定，规模随输入目标与资源数量变化。",
    )
    add_equation(document, "pₖ⁻(c) = Σᵣ P(c｜r，Δt) · pₖ₋₁⁺(r)")
    add_body(
        document,
        "式中，pₖ₋₁⁺(r)表示上一周期网格r的后验概率，P(c｜r，Δt)表示经时间Δt由r转移至c的概率，pₖ⁻(c)表示当前先验概率。该式用于按运动模型和信息时延传播目标存在概率。转移分布随目标机动程度和时间误差扩大，预测中心只作为搜索依据。",
    )
    add_body(
        document,
        "先验概率形成后，依据相机位姿、云台角度、视场、遮挡、预计像素尺寸和检测条件计算实际可观测区域。有效发现按观测方向与量测误差形成相符程度，未发现按真实覆盖区域的预计探测概率形成负证据。更新过程保存原始概率、观测时刻、原因和实际视锥，使概率变化可追溯。",
    )
    add_equation(document, "pₖ⁺(c) = pₖ⁻(c) · L(zₖ｜c) ／ Σᵣ[pₖ⁻(r) · L(zₖ｜r)]")
    add_body(
        document,
        "式中，pₖ⁻(c)表示先验概率，L(zₖ｜c)表示网格c产生观测zₖ的似然，分母完成归一化，pₖ⁺(c)表示后验概率。该式用于将发现或未发现证据纳入概率分布。有效发现提高观测一致区域概率；未发现只降低真实视锥内具备探测条件的区域概率，遮挡区和像素不足区不参与负更新。",
    )
    add_body(
        document,
        "图2给出概率更新与发现后任务组织的关系。初始概率反映粗略位置和运动传播，未发现后已观察区域概率降低，连续发现后概率向观测一致区域集中。分配阶段只保留通过时效、可达、云台、通信和安全检查的候选关系，由确定性求解器形成一一任务结果。",
    )
    add_figure(document, FIGURES[1], "动态概率搜索、发现与稀疏一一分配原理", 2)

    add_body(
        document,
        "随后生成候选观察动作，每个动作由平台观察点、到达方向、预计到达时刻、云台方位与俯仰以及观察持续时间共同组成。满足空域、速度、转弯、云台角限、成像尺寸、通信和安全间隔要求的动作方可进入收益比较。固定规则先形成可独立运行的搜索基线，并记录每个动作的收益分项、约束余量和拒绝原因。",
    )
    add_equation(document, "Uᵢc = wᵢ I(c) − wₜ Tᵢc + wɢ Gᵢc − wᵣ Rᵢc − w꜀ Cᵢc")
    add_body(
        document,
        "式中，Uᵢc表示无人机i执行观察动作c的综合收益，I(c)表示预计信息收益，Tᵢc表示飞行、转向和稳定后的总到达时间，Gᵢc表示云台可达余量，Rᵢc与Cᵢc分别表示重复观察和航路冲突代价，各权重表示对应分项的重要程度。该式用于对全部可行观察动作进行统一排序。各分项采用统一时间基准和可核对单位，权重在正式验证前确定并进行灵敏度检查。",
    )
    add_body(
        document,
        "规则基线依据收益矩阵安排观察单元，并记录资源不足、预计超时或安全冲突原因。强化学习只读取压缩概率、近期观察、资源状态和规则收益，输出候选优先级修正。训练评价采用发现率、首次发现时间、重复覆盖和计划稳定性，在线输入不含真实身份，模型不得绕过规则候选集合。",
    )
    add_equation(document, "Ũᵢc = Uᵢc + β · tanh(Δuᵢc)，0 ≤ β ≤ βₘₐₓ")
    add_body(
        document,
        "式中，Ũᵢc表示修正后的观察收益，Uᵢc表示规则基线收益，Δuᵢc表示学习方法给出的排序增量，β表示修正幅度，βₘₐₓ表示批准上限，tanh函数将异常增量压缩至有限范围。该式用于把学习输出限制为规则收益附近的有界优先级修正。禁飞、不可达和云台受限动作仍由确定性约束剔除；模型不可用、输入越界或推理超时时，系统直接采用规则基线。",
    )
    add_body(
        document,
        "确定性求解器依据排序形成平台方向和云台指向候选，安全检查通过后由控制器执行。实际到位时间、云台角度、真实视锥、取帧时刻和观察结果全部回写，概率更新只使用执行记录。有效未发现、连续发现、资源释放或高概率区域转移时启动下一轮计算，形成联合滚动搜索。",
    )

    add_heading(document, "（二）发现后的动态分配与安全重算", level=2)
    add_body(
        document,
        "连续发现后先形成待确认的局部航迹，记录位置或视线估计、协方差、双时间戳、来源平台和证据成熟度。可用无人机与待分目标构成关系图两侧，只为任务窗口内可形成有效观察的组合建立候选边。稀疏候选图减少无效比较，并在学习计算前固定时效、可达、云台、任务、通信和安全条件。",
    )
    add_equation(document, "gᵢⱼ = g时效 · g可达 · g云台 · g任务 · g通信 · g安全 ∈ {0，1}")
    add_body(
        document,
        "式中，gᵢⱼ表示无人机i与目标j的准入结果，六个因子分别表示时效、平台可达、云台可达、任务占用、通信和安全条件，任一为0即剔除该候选边。该式用于在代价计算前形成可审计的稀疏无人机—目标图。各项检查均保存输入、阈值和拒绝原因，高优先目标也须满足全部硬约束。",
    )
    add_body(
        document,
        "通过准入的候选边进入规则代价计算。代价包括预计到达时间、平台与云台机动难度、目标证据风险、当前视觉连续性、航路冲突、资源余量和换令影响，并以目标紧迫程度形成收益项。规则代价必须具备独立完成分配的能力，学习方法仅在多个可行组合代价接近或需要考虑多周期资源价值时提供限定幅度的辅助修正。",
    )
    add_equation(document, "C规则(i，j) = aTᵢⱼ + bMᵢⱼ + cQⱼ + dVᵢⱼ + eHᵢⱼ + fEᵢ + sSᵢⱼ − hWⱼ")
    add_body(
        document,
        "式中，C规则(i，j)表示规则代价，T表示到达时间，M表示机动难度，Q表示证据风险，V表示视觉连续性，H表示航路冲突，E表示资源余量，S表示换令影响，W表示目标紧迫程度，a、b、c、d、e、f、s和h为权重。该式用于比较全部准入候选的综合任务代价。各分项保存原始量、归一化量和权重，目标紧迫程度只改变可行候选顺序。",
    )
    add_body(
        document,
        "图网络利用候选边的竞争关系评估资源与目标适配程度，强化学习利用近期结果估计任务保持与后续资源需求。两类方法只输出候选代价修正，并经过数值压缩、幅度限制、输入范围检查和版本核验。模型不得新增候选边、直接确认目标身份、指定负责无人机或生成执行计划。",
    )
    add_equation(document, "C最终(i，j) = C规则(i，j) + α · tanh(ΔCᵢⱼ)，0 ≤ α ≤ αₘₐₓ")
    add_body(
        document,
        "式中，C最终(i，j)表示进入确定性求解器的候选代价，C规则(i，j)表示规则基线代价，ΔCᵢⱼ表示学习方法给出的代价增量，α表示修正幅度，αₘₐₓ表示批准上限。该式用于将图网络或强化学习输出压缩为规则代价附近的有界修正。修正比例、输入范围、模型版本或推理时延不符合要求时，本轮修正失效并恢复规则代价，候选筛选条件保持不变。",
    )
    add_body(
        document,
        "最终任务关系采用匈牙利算法或最小费用流求解。单周期一一关系使用匈牙利算法，需要显式表示未分配、连续任务或资源容量时使用最小费用流；两种方法均保证一架无人机同一时段最多负责一个主要目标，一个目标最多由一架主要无人机负责。求解器允许目标暂时未分配，并记录对应代价和原因，避免为满足矩阵完整性而发布不可行任务。",
    )
    add_equation(document, "min ΣᵢΣⱼ xᵢⱼ C最终(i，j) + Σⱼ uⱼ C未分(j)；Σⱼxᵢⱼ≤1，Σᵢxᵢⱼ+uⱼ=1")
    add_body(
        document,
        "式中，xᵢⱼ表示无人机i是否负责目标j，uⱼ表示目标j是否未分配，C最终(i，j)表示候选代价，C未分(j)表示未分配代价；两个约束限制单机任务数量，并保证目标获得唯一主要资源或保持未分配。该式用于在可行候选中最小化总体任务代价。最终关系由确定性求解结果发布，学习分数只参与代价计算。",
    )
    add_body(
        document,
        "任务发布后进入保持机制，并采用事件触发方式处理重大变化。普通代价波动只有在新方案改善超过门槛且现有关系达到最短保持时间时才允许换令；高优先目标新增、目标失效、资源故障、任务不可达、身份冲突、航路冲突和计划过期等硬事件直接触发重算。重算优先保留身份稳定、仍可达且处于有效期内的关系，只对新增、失效和冲突部分重新组织。",
    )
    add_equation(document, "J新 ≤ J现 − ΔJ，且 t − t上次 ≥ T保持")
    add_body(
        document,
        "式中，J新与J现分别表示新方案和现行方案的总代价，ΔJ表示允许换令所需的最小改善量，t−t上次表示现行任务已保持的时间，T保持表示最短保持时间。该式用于普通代价波动条件下判断是否启动重分配。硬事件可直接触发重算，但新计划仍须通过版本、有效期、资源唯一性和安全检查。",
    )
    add_body(
        document,
        "计划包含发布者、递增版本、生效与失效时间、任务关系和触发原因，执行端拒绝旧版本及过期计划。发布前复查禁入区、航路冲突、机间间隔、云台、通信、能源、目标证据和资源唯一性，失败候选移除后重新求解。计算超时或输入异常时保持安全有效的现行任务，并以规则基线处理受影响部分，学习输出不得越过安全检查。",
    )
    add_body(
        document,
        "图3给出从粗略线索到任务关系的完整流程。概率更新与滚动搜索形成观察和更新闭环，连续发现后进入稀疏分配流程。权限边界规定学习方法只调整搜索优先级或候选代价，平台控制、云台控制、一一求解、计划发布和安全检查由确定性链路承担。",
    )
    add_figure(document, FIGURES[2], "不确定信息驱动的协同搜索与动态分配完整算法流程", 3)

    add_heading(document, "三、期望结果与后续实施", level=1, page_break_before=True)
    add_heading(document, "（一）现有基础、待验证环节与考核目标", level=2)
    add_body(
        document,
        "截至2026年8月，已完成适应不同目标与资源数量的概率搜索单元、规则收益矩阵、匈牙利唯一分配、真实视锥负观测、不使用真实目标编号的在线记录和连续两帧确认等基础实现。离线矩阵覆盖20目标/8机、20目标/30机和40目标/50机，设置30米、60米、100米误差，每档5个固定随机场景，共45组；20目标/8机三档平均连续确认率为100%、97%和91%，其余规模均为100%。该回放在0.8秒后采用线性方法预测目标位置，并使用理想化速度、云台和探测条件，未加入随机漏检、虚警、导航误差、防碰撞及通信时延，属于科研仿真证据。",
    )
    add_body(
        document,
        "AirSim接口验证中，20目标/8机场景按连续3帧确认条件重新试验，20个目标均达到10像素识别门限，其中19个完成连续确认。20目标/30机和40目标/50机的单一固定随机场景分别确认20/20和40/40，规划平均耗时为11.739毫秒和35.753毫秒。后续仍需补充多个固定随机场景、完整18秒轨迹、平台加速度、云台稳定、导航姿态误差、网络时延、漏检虚警和防碰撞复核，现有结果不代表设备性能或规模上限。",
    )
    add_body(
        document,
        "学习修正处于离线研究和并行对照阶段，模型输出只记录、不参与任务执行，尚无在线分配、计划发布或控制权限。输入限定为不含真实目标身份的态势和规则候选，输出只修正搜索优先级或已准入候选代价，并经幅度限制、异常检查和确定性安全约束裁剪；确定性求解器、版本检查和独立安全检查作出最终决定。后续试验预先确定场景、规模、时间预算、传感器条件、随机场景、规则基线、统计方法和判定门槛，同时报告均值、最差场景与置信区间。",
    )
    add_table_title(document, 1, "建议考核指标及验收方法")
    add_table(
        document,
        ["指标", "计算方法", "建议验收判据", "数据来源"],
        [
            ["限时连续发现率", "时间预算内形成连续局部航迹的目标数／目标总数", "同场景配对结果不低于规则基线，并报告最差场景", "相机取帧、视锥和局部航迹记录"],
            ["首次发现时间", "从搜索开始到首次有效发现的中位数、95%分位和超时率", "在发现率不下降前提下不长于规则基线", "统一时钟事件日志"],
            ["概率校准误差", "以独立真值计算布里尔分数或分箱校准误差", "不高于规则概率更新基线，训练集与测试集分开报告", "概率网格快照与离线真值"],
            ["单位时间信息收益", "实际概率熵下降量／搜索用时，并分解发现与未发现贡献", "配对均值不低于规则基线", "更新前后概率快照"],
            ["重复搜索与实际覆盖", "重复观察概率质量占比、真实视锥覆盖概率质量和无效到位次数", "发现率不下降时重复搜索不高于基线", "平台位姿、云台角和视锥记录"],
            ["滚动计算时延", "搜索排序、候选图、确定性求解和安全检查的95%与99%分位", "99%分位不超过滚动周期的一半，周期值待设备标定", "单调时钟分段日志"],
            ["分配可行性与稳定性", "可行计划率、任务保持时长、换令次数和未分配原因占比", "可行计划率不低于基线，非硬事件换令不高于基线", "版本化计划与执行回执"],
            ["安全与计划一致性", "航路冲突、重复主要任务、不可达发布、过期计划执行和旧版本接收次数", "各项均为0", "安全检查、计划和控制回执"],
            ["学习权限边界", "绕过准入检查、超限修正、直接产生平台或云台控制、直接发布任务的次数", "各项均为0", "模型输入输出与安全裁剪记录"],
            ["跨规模泛化", "未参与训练的2对2、4对4、6对6场景中上述指标的配对差值", "每级独立报告且不以总体平均掩盖单级退化", "预先确定的场景、随机场景和独立评分结果"],
        ],
        [2.8, 5.0, 4.4, 3.5],
    )
    add_body(
        document,
        "每项指标必须同时给出单位、测试条件、样本数、统计方法和原始数据来源。发现率提高但重复搜索、换令或安全事件增加时，不判定为总体改善；某一规模明显退化时，单独列出该规模结果。学习方法只有在权限边界零违规、关键效果不低于规则基线且计算时延满足滚动周期要求时，才可进入下一阶段验证。",
    )

    add_heading(document, "（二）分阶段实施路线", level=2)
    add_body(
        document,
        "实施路线按照质点快速训练、AirSim验证、2对2至6对6封闭实验室验证和后续扩规模四个阶段推进。每个阶段先运行完整规则基线，再开展学习方法并行对照，达到阶段进入条件后方可在批准幅度内启用优先级或候选代价修正。出现安全违规、数据记录不完整、结果无法重放或效果低于预先确定的门槛时，恢复规则基线并定位原因，当前阶段通过后再扩大试验规模。",
    )
    add_figure(document, FIGURES[3], "从质点快速训练到封闭实验室验证及后续扩规模路线", 4)
    add_table_title(document, 2, "分阶段工作、量测输出与进入条件")
    add_table(
        document,
        ["阶段", "核心工作", "可量测输出", "进入下一阶段条件"],
        [
            ["质点快速训练", "批量生成目标机动、位置误差、资源不足、视场变化、故障和通信退化场景；完成规则基线、优先级学习和候选代价修正的分项对照", "发现率、首次发现时间、概率校准、重复搜索、分配代价、换令次数、计算时延和安全事件", "训练、验证、测试严格隔离；规则链路独立可用；安全违规为0；独立测试不低于预先确定的基线"],
            ["AirSim验证", "接入相机视锥、平台运动、云台角限与稳定时间，验证未发现更新、到位取帧、航路冲突、事件重算和端到端时延", "实际视锥覆盖、运动与云台跟踪误差、搜索与分配全过程、时延分解、日志完整率和重放一致性", "全过程数据可追溯并可重放；安全违规为0；关键指标达到预先确定的门槛"],
            ["封闭实验室验证", "按2对2、4对4、6对6逐级开展封闭条件试验，统一标定相机、云台、定位、网络和时钟，每级使用固定场景进行规则与学习配对复验", "各规模发现、分配、任务保持、冲突规避、故障恢复、时延和最差重复场景结果", "每一级单独通过后再扩到下一级；原始记录、标定文件、版本和独立评分齐全"],
            ["后续扩规模", "在6对6达到门槛后逐步增加目标与资源数量，加入遮挡、密集交叉、资源失效和通信受限，形成规模与实时性的边界曲线", "规模—效果—时延曲线、资源利用率、退化模式、未见场景泛化和经验证的适用范围", "不预设可扩上限；每次扩展均沿用同一安全门、配对基线和证据要求"],
        ],
        [2.6, 5.2, 4.2, 3.7],
    )
    add_body(
        document,
        "各阶段均形成可复核的指标证据。场景、随机种子、输入摘要、概率网格、候选动作、学习修正、安全裁剪、最终计划、执行回执和评分结果绑定同一版本保存。阶段评价分别回答有限视场下能否及时缩小目标可能区域、发现后能否形成稳定唯一任务关系，以及学习辅助在确定性安全边界内能否产生可重复的增益。",
    )

    properties = document.core_properties
    properties.title = TITLE
    properties.subject = "有限视场协同搜索、不确定信息更新、学习增强动态分配与分阶段验证"
    properties.author = "项目组"
    properties.keywords = "有限视场, 协同搜索, 概率网格, 动态分配, 匈牙利算法, 最小费用流, 安全边界"
    properties.comments = "Word-only leadership scientific-problem report"
    document.save(OUTPUT)


def document_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def body_text(document: Document) -> str:
    parts: list[str] = []
    started = False
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "一、任务、主要难点与关键技术":
            started = True
        if started:
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def validate_assets() -> list[dict[str, object]]:
    results = []
    for path in FIGURES:
        if not path.exists():
            raise RuntimeError(f"missing figure: {path}")
        with Image.open(path) as image:
            width, height = image.size
            dpi = image.info.get("dpi", (0, 0))
            if image.format != "PNG":
                raise RuntimeError(f"figure is not PNG: {path}")
            if width < 3000 or height < 1700:
                raise RuntimeError(f"figure resolution too low: {path.name}={width}x{height}")
            if min(dpi) < 295:
                raise RuntimeError(f"figure DPI too low: {path.name}={dpi}")
            results.append({"name": path.name, "width": width, "height": height, "dpi": tuple(round(v, 1) for v in dpi)})
    return results


def validate_document() -> dict[str, object]:
    document = Document(OUTPUT)
    with ZipFile(OUTPUT) as archive:
        damaged = archive.testzip()
        if damaged is not None:
            raise RuntimeError(f"damaged DOCX member: {damaged}")
        media = [name for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/")]
    if len(media) != 4 or len(document.inline_shapes) != 4:
        raise RuntimeError(f"expected 4 embedded figures, found media={len(media)}, inline={len(document.inline_shapes)}")
    if len(document.tables) != 2:
        raise RuntimeError(f"expected 2 tables, found {len(document.tables)}")

    chapters = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"]
    expected_chapters = [
        "一、任务、主要难点与关键技术",
        "二、具体算法步骤与实施流程",
        "三、期望结果与后续实施",
    ]
    if chapters != expected_chapters:
        raise RuntimeError(f"unexpected chapter structure: {chapters}")

    full_text = document_text(document)
    report_body = body_text(document)
    cjk_chars = len(CJK_RE.findall(report_body))
    if not 4000 <= cjk_chars <= 6000:
        raise RuntimeError(f"body CJK character count out of range: {cjk_chars}")

    sensitive_counts = {phrase: full_text.count(phrase) for phrase in SENSITIVE_PHRASES}
    remaining_sensitive = {phrase: count for phrase, count in sensitive_counts.items() if count}
    if remaining_sensitive:
        raise RuntimeError(f"sensitive wording remains: {remaining_sensitive}")

    required_terms = (
        KEY_TECHNOLOGY,
        "动态概率网格",
        "未发现",
        "信息收益",
        "到达时间",
        "云台可达",
        "重复搜索",
        "航路冲突",
        "规则基线",
        "强化学习",
        "稀疏无人机—目标图",
        "规则代价",
        "图网络",
        "有界修正",
        "匈牙利算法",
        "最小费用流",
        "一一任务关系",
        "任务保持",
        "事件触发",
        "安全检查",
        "质点快速训练",
        "AirSim验证",
        "2对2",
        "4对4",
        "6对6",
        "后续扩规模",
    )
    for term in required_terms:
        if term not in full_text:
            raise RuntimeError(f"missing required term: {term}")
    for forbidden in ("合同", "槽位", "模块编号", "开发记录"):
        if forbidden in full_text:
            raise RuntimeError(f"forbidden term found: {forbidden}")

    picture_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph._p.xpath(".//w:drawing")]
    if len(picture_paragraphs) != 4:
        raise RuntimeError(f"expected 4 picture paragraphs, found {len(picture_paragraphs)}")
    for paragraph in picture_paragraphs:
        if paragraph.paragraph_format.line_spacing_rule != WD_LINE_SPACING.SINGLE:
            raise RuntimeError("picture paragraph must use single line spacing")
        if paragraph.paragraph_format.line_spacing != 1.0:
            raise RuntimeError("picture paragraph must use non-fixed 1.0 line spacing")

    body_started = False
    narrative_paragraphs = []
    formula_count = 0
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text == expected_chapters[0]:
            body_started = True
        if not body_started or not text:
            continue
        if paragraph.style.name == "Normal":
            narrative_paragraphs.append(text)
            if len(SENTENCE_END_RE.findall(text)) < 3:
                raise RuntimeError(f"body paragraph has fewer than 3 sentences: {text[:60]}")
        if paragraph.style.name == "公式":
            formula_count += 1
            next_index = index + 1
            while next_index < len(document.paragraphs) and not document.paragraphs[next_index].text.strip():
                next_index += 1
            if next_index >= len(document.paragraphs):
                raise RuntimeError(f"formula lacks immediate explanation: {text}")
            explanation = document.paragraphs[next_index].text
            if not explanation.startswith("式中") or "该式用于" not in explanation:
                raise RuntimeError(f"formula lacks formal variable and purpose explanation: {text}")
    if formula_count != 9:
        raise RuntimeError(f"expected 9 formulas, found {formula_count}")

    return {
        "paragraphs": len(document.paragraphs),
        "narrative_paragraphs": len(narrative_paragraphs),
        "tables": len(document.tables),
        "images": len(media),
        "formulas": formula_count,
        "cjk_chars": cjk_chars,
        "sensitive_phrases": sum(sensitive_counts.values()),
        "bytes": OUTPUT.stat().st_size,
    }


def main() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    build_figures()
    asset_metrics = validate_assets()
    build_document()
    document_metrics = validate_document()
    print(f"output={OUTPUT}")
    print(
        "document="
        f"paragraphs:{document_metrics['paragraphs']}, "
        f"narrative:{document_metrics['narrative_paragraphs']}, "
        f"tables:{document_metrics['tables']}, images:{document_metrics['images']}, "
        f"formulas:{document_metrics['formulas']}, cjk_chars:{document_metrics['cjk_chars']}, "
        f"sensitive_phrases:{document_metrics['sensitive_phrases']}, bytes:{document_metrics['bytes']}"
    )
    for item in asset_metrics:
        print(f"figure={item['name']} {item['width']}x{item['height']} dpi={item['dpi']}")


if __name__ == "__main__":
    main()
