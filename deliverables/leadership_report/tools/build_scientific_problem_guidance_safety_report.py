#!/usr/bin/env python3
"""Build the Word-only leadership report and its four Chinese figures."""

from __future__ import annotations

import math
import re
from pathlib import Path
from zipfile import ZipFile

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from matplotlib.patches import Circle, FancyArrowPatch, FancyBboxPatch, Polygon, Rectangle, Wedge
from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
REPORT_DIR = ROOT / "deliverables" / "leadership_report"
ASSET_DIR = REPORT_DIR / "assets" / "scientific_problem_guidance_safety"
OUTPUT_DOCX = REPORT_DIR / "多机末端接近制导与防碰撞统一科学问题报告_CN.docx"

FIGURES = {
    "crossing": ASSET_DIR / "01_多航迹交叉三维场景.png",
    "correction": ASSET_DIR / "02_正常制导与安全修正原理.png",
    "flow": ASSET_DIR / "03_完整控制流程图.png",
    "route": ASSET_DIR / "04_封闭实验室实施路线.png",
}

TITLE = "多机末端接近制导与防碰撞统一科学问题报告"
SHORT_TITLE = "多机末端接近制导与防碰撞统一科学问题"
BODY_FONT = "仿宋"
HEADING_FONT = "黑体"
ASCII_FONT = "Times New Roman"
MATH_FONT = "Cambria Math"
FIGURE_FONT_FILE = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")

NAVY = "#264653"
TEAL = "#2A9D8F"
GOLD = "#E9C46A"
ORANGE = "#F4A261"
RED = "#C94C4C"
BLUE = "#3E6EA8"
PURPLE = "#745A9C"
GREEN = "#4F8A5B"
LIGHT = "#F5F7F8"
MID = "#D7DEE2"
INK = "#1F2933"
MUTED = "#5C6770"


def configure_matplotlib() -> font_manager.FontProperties:
    if not FIGURE_FONT_FILE.exists():
        raise FileNotFoundError(f"Chinese figure font not found: {FIGURE_FONT_FILE}")
    font_manager.fontManager.addfont(str(FIGURE_FONT_FILE))
    prop = font_manager.FontProperties(fname=str(FIGURE_FONT_FILE))
    plt.rcParams.update(
        {
            "font.family": prop.get_name(),
            "axes.unicode_minus": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
            "axes.edgecolor": MID,
            "axes.labelcolor": INK,
            "xtick.color": MUTED,
            "ytick.color": MUTED,
        }
    )
    return prop


def quadratic_bezier(start: np.ndarray, control: np.ndarray, end: np.ndarray, count: int = 140) -> np.ndarray:
    t = np.linspace(0.0, 1.0, count)[:, None]
    return (1.0 - t) ** 2 * start + 2.0 * (1.0 - t) * t * control + t**2 * end


def draw_crossing_scene(font_prop: font_manager.FontProperties) -> None:
    fig, ax = plt.subplots(figsize=(12.0, 7.5), dpi=300)
    ax.set_xlim(-130, 125)
    ax.set_ylim(-72, 75)
    ax.set_aspect("equal")
    ax.axis("off")

    def project(points: np.ndarray) -> np.ndarray:
        values = np.atleast_2d(points).astype(float)
        screen_x = values[:, 0] - 0.58 * values[:, 1]
        screen_y = 0.30 * values[:, 0] + 0.22 * values[:, 1] + 1.75 * values[:, 2] - 56.0
        return np.column_stack([screen_x, screen_y])

    grid_color = "#D8DEE2"
    for x in np.arange(-90, 91, 30):
        line = project(np.array([[x, -75, 10], [x, 75, 10]]))
        ax.plot(line[:, 0], line[:, 1], color=grid_color, linewidth=0.8, zorder=0)
    for y in np.arange(-75, 76, 25):
        line = project(np.array([[-95, y, 10], [95, y, 10]]))
        ax.plot(line[:, 0], line[:, 1], color=grid_color, linewidth=0.8, zorder=0)

    origin = project(np.array([[0, 0, 10]]))[0]
    axis_n = project(np.array([[85, 0, 10]]))[0]
    axis_e = project(np.array([[0, 70, 10]]))[0]
    axis_h = project(np.array([[0, 0, 62]]))[0]
    for end, color in ((axis_n, NAVY), (axis_e, NAVY), (axis_h, NAVY)):
        ax.add_patch(FancyArrowPatch(tuple(origin), tuple(end), arrowstyle="-|>", mutation_scale=13, color=color, linewidth=1.4, zorder=1))
    ax.text(axis_n[0] + 2, axis_n[1], "北向", fontsize=10.5, color=NAVY, fontproperties=font_prop)
    ax.text(axis_e[0] - 14, axis_e[1] + 1, "东向", fontsize=10.5, color=NAVY, fontproperties=font_prop)
    ax.text(axis_h[0] + 2, axis_h[1] + 1, "高度", fontsize=10.5, color=NAVY, fontproperties=font_prop)

    target_specs = [
        (np.array([10.0, -62.0, 32.0]), np.array([23.0, 2.0, 34.0]), np.array([68.0, 58.0, 39.0]), RED, "目标一"),
        (np.array([75.0, -48.0, 46.0]), np.array([18.0, -2.0, 37.0]), np.array([-54.0, 52.0, 31.0]), ORANGE, "目标二"),
        (np.array([-68.0, 28.0, 26.0]), np.array([2.0, 7.0, 35.0]), np.array([76.0, -22.0, 43.0]), PURPLE, "目标三"),
    ]
    interceptor_specs = [
        (np.array([-84.0, -58.0, 18.0]), np.array([-25.0, -17.0, 29.0]), np.array([27.0, 5.0, 34.0]), BLUE, "拦截机一"),
        (np.array([88.0, 55.0, 20.0]), np.array([45.0, 18.0, 32.0]), np.array([21.0, 0.0, 37.0]), TEAL, "拦截机二"),
        (np.array([-72.0, 66.0, 52.0]), np.array([-30.0, 30.0, 43.0]), np.array([7.0, 8.0, 36.0]), GREEN, "拦截机三"),
    ]

    for start, control, end, color, label in target_specs:
        path = quadratic_bezier(start, control, end)
        screen = project(path)
        ground = project(np.column_stack([path[:, 0], path[:, 1], np.full(len(path), 10.0)]))
        ax.plot(ground[:, 0], ground[:, 1], color=color, linewidth=1.1, linestyle="--", alpha=0.35, zorder=1)
        ax.plot(screen[:, 0], screen[:, 1], color=color, linewidth=3.1, label=label, zorder=4)
        start_screen = project(start)[0]
        end_screen = project(end)[0]
        ax.scatter(*start_screen, color=color, s=65, marker="^", edgecolor="white", linewidth=0.8, zorder=6)
        ax.scatter(*end_screen, color=color, s=35, marker="o", edgecolor="white", linewidth=0.8, zorder=6)
        stem = project(np.array([[start[0], start[1], 10.0], start]))
        ax.plot(stem[:, 0], stem[:, 1], color=color, linewidth=0.8, linestyle=":", alpha=0.65, zorder=2)

    for start, control, end, color, label in interceptor_specs:
        path = quadratic_bezier(start, control, end)
        screen = project(path)
        ground = project(np.column_stack([path[:, 0], path[:, 1], np.full(len(path), 10.0)]))
        ax.plot(ground[:, 0], ground[:, 1], color=color, linewidth=1.1, linestyle="--", alpha=0.35, zorder=1)
        ax.plot(screen[:, 0], screen[:, 1], color=color, linewidth=3.3, label=label, zorder=5)
        start_screen = project(start)[0]
        end_screen = project(end)[0]
        ax.scatter(*start_screen, color=color, s=70, marker="s", edgecolor="white", linewidth=0.8, zorder=6)
        ax.scatter(*end_screen, color=color, s=38, marker="o", edgecolor="white", linewidth=0.8, zorder=6)
        extension = np.vstack([end, end + (end - path[-18]) * 1.4])
        extension_screen = project(extension)
        ax.plot(extension_screen[:, 0], extension_screen[:, 1], color=color, linestyle="--", linewidth=1.8, alpha=0.8, zorder=4)
        stem = project(np.array([[start[0], start[1], 10.0], start]))
        ax.plot(stem[:, 0], stem[:, 1], color=color, linewidth=0.8, linestyle=":", alpha=0.65, zorder=2)

    center = np.array([17.0, 4.0, 36.0])
    center_screen = project(center)[0]
    ax.add_patch(Circle(tuple(center_screen), 13.0, facecolor="#FDECEC", edgecolor=RED, linewidth=1.8, alpha=0.65, zorder=2))
    ax.add_patch(Circle(tuple(center_screen), 8.0, facecolor="none", edgecolor=RED, linewidth=0.9, linestyle="--", alpha=0.8, zorder=3))
    ax.text(center_screen[0] + 10, center_screen[1] + 14, "航迹汇聚风险区", color=RED, fontsize=12.5, fontproperties=font_prop, weight="bold", zorder=7)

    fov_origin = interceptor_specs[0][0]
    fov_end = fov_origin + np.array([35.0, 23.0, 13.0])
    fov_origin_screen = project(fov_origin)[0]
    for offset in (-7.5, 7.5):
        fov_tip = np.array([fov_end[0], fov_end[1] + offset, fov_end[2] + offset * 0.25])
        fov_tip_screen = project(fov_tip)[0]
        ax.plot([fov_origin_screen[0], fov_tip_screen[0]], [fov_origin_screen[1], fov_tip_screen[1]], color=GOLD, linewidth=1.8, alpha=0.9, zorder=3)
    label_point = project(fov_origin + np.array([18.0, 8.0, 12.0]))[0]
    ax.text(label_point[0], label_point[1], "有限视场", color="#8A681A", fontsize=11.5, fontproperties=font_prop, zorder=7)

    ax.set_title("多机分别追踪时的三维航迹交叉与汇聚", fontsize=20, color=INK, pad=18, fontproperties=font_prop)
    ax.text(-145, -37, "虚线为地面投影，竖向点线表示航迹高度", fontsize=10.5, color=MUTED, fontproperties=font_prop)
    handles, labels = ax.get_legend_handles_labels()
    ax.legend(
        handles,
        labels,
        loc="upper left",
        bbox_to_anchor=(0.005, 0.98),
        ncol=2,
        frameon=True,
        framealpha=0.96,
        prop=font_manager.FontProperties(fname=str(FIGURE_FONT_FILE), size=10.5),
    )
    fig.text(
        0.5,
        0.025,
        "每架拦截机只追踪自身目标时，正常指令仍可能把多条航迹带入同一狭小空域。",
        ha="center",
        va="bottom",
        color=MUTED,
        fontsize=11.5,
        fontproperties=font_prop,
    )
    fig.subplots_adjust(left=0.01, right=0.99, bottom=0.08, top=0.91)
    fig.savefig(FIGURES["crossing"], dpi=300, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def draw_nominal_and_correction(font_prop: font_manager.FontProperties) -> None:
    fig, axes = plt.subplots(1, 2, figsize=(12.0, 7.2), dpi=300)
    fig.suptitle("正常拦截指令与风险触发后的最小安全修正", fontsize=20, color=INK, y=0.96, fontproperties=font_prop)

    for ax in axes:
        ax.set_xlim(0, 10)
        ax.set_ylim(0, 8)
        ax.set_aspect("equal")
        ax.axis("off")

    left = axes[0]
    left.add_patch(FancyBboxPatch((0.25, 0.25), 9.5, 7.2, boxstyle="round,pad=0.02,rounding_size=0.08", facecolor=LIGHT, edgecolor=MID, linewidth=1.4))
    left.text(5, 7.05, "无预测冲突：保留正常制导", ha="center", va="center", fontsize=15, color=NAVY, fontproperties=font_prop, weight="bold")
    left.add_patch(Wedge((2.0, 2.0), 3.2, 8, 58, facecolor=GOLD, alpha=0.20, edgecolor=GOLD, linewidth=1.2))
    left.add_patch(Rectangle((1.65, 1.65), 0.7, 0.7, facecolor=BLUE, edgecolor="white", linewidth=1.0))
    left.add_patch(Rectangle((1.65, 5.25), 0.7, 0.7, facecolor=TEAL, edgecolor="white", linewidth=1.0))
    left.add_patch(Polygon([[8.0, 3.5], [7.55, 3.2], [7.55, 3.8]], closed=True, facecolor=RED, edgecolor="white"))
    left.add_patch(Polygon([[8.1, 6.3], [7.65, 6.0], [7.65, 6.6]], closed=True, facecolor=ORANGE, edgecolor="white"))
    left.annotate("", xy=(7.35, 3.5), xytext=(2.35, 2.0), arrowprops=dict(arrowstyle="->", lw=3.2, color=BLUE))
    left.annotate("", xy=(7.45, 6.25), xytext=(2.35, 5.6), arrowprops=dict(arrowstyle="->", lw=3.2, color=TEAL))
    left.text(4.8, 2.45, "正常制导指令", color=BLUE, fontsize=12, fontproperties=font_prop)
    left.text(4.6, 5.85, "正常制导指令", color=TEAL, fontsize=12, fontproperties=font_prop)
    left.text(2.0, 1.25, "目标位于视场内", ha="center", color="#806117", fontsize=11.5, fontproperties=font_prop)
    left.text(5.0, 0.72, "未触发安全层，实际指令等于原指令", ha="center", color=MUTED, fontsize=11.5, fontproperties=font_prop)

    right = axes[1]
    right.add_patch(FancyBboxPatch((0.25, 0.25), 9.5, 7.2, boxstyle="round,pad=0.02,rounding_size=0.08", facecolor=LIGHT, edgecolor=MID, linewidth=1.4))
    right.text(5, 7.05, "预测有冲突：安全层最小改动", ha="center", va="center", fontsize=15, color=RED, fontproperties=font_prop, weight="bold")
    starts = [(1.5, 1.4, BLUE), (1.4, 6.1, TEAL)]
    goals = [(8.4, 6.1, RED), (8.5, 1.4, ORANGE)]
    for x, y, c in starts:
        right.add_patch(Rectangle((x - 0.33, y - 0.33), 0.66, 0.66, facecolor=c, edgecolor="white"))
    for x, y, c in goals:
        right.add_patch(Polygon([[x + 0.35, y], [x - 0.25, y - 0.34], [x - 0.25, y + 0.34]], closed=True, facecolor=c, edgecolor="white"))
    right.plot([1.8, 8.0], [1.7, 5.8], linestyle="--", linewidth=2.2, color=BLUE, alpha=0.7)
    right.plot([1.7, 8.1], [5.8, 1.7], linestyle="--", linewidth=2.2, color=TEAL, alpha=0.7)
    right.add_patch(Circle((4.9, 3.75), 0.92, facecolor="#FDECEC", edgecolor=RED, linewidth=2.0, alpha=0.95))
    right.text(4.9, 3.75, "预计\n冲突区", ha="center", va="center", color=RED, fontsize=11.5, fontproperties=font_prop, weight="bold")
    x1 = np.linspace(1.8, 8.0, 100)
    y1 = 1.7 + 4.1 * (x1 - 1.8) / 6.2 - 1.35 * np.exp(-((x1 - 4.8) / 1.2) ** 2)
    x2 = np.linspace(1.7, 8.1, 100)
    y2 = 5.8 - 4.1 * (x2 - 1.7) / 6.4 + 1.35 * np.exp(-((x2 - 5.0) / 1.2) ** 2)
    right.plot(x1, y1, color=BLUE, linewidth=3.4)
    right.plot(x2, y2, color=TEAL, linewidth=3.4)
    right.annotate("", xy=(7.95, y1[-1]), xytext=(7.35, y1[-8]), arrowprops=dict(arrowstyle="->", lw=2.8, color=BLUE))
    right.annotate("", xy=(8.05, y2[-1]), xytext=(7.45, y2[-8]), arrowprops=dict(arrowstyle="->", lw=2.8, color=TEAL))
    right.annotate("安全修正", xy=(4.2, 2.9), xytext=(2.6, 2.1), fontsize=11.5, color=RED, fontproperties=font_prop, arrowprops=dict(arrowstyle="->", color=RED, lw=1.5))
    right.annotate("安全修正", xy=(5.7, 4.65), xytext=(6.4, 5.55), fontsize=11.5, color=RED, fontproperties=font_prop, arrowprops=dict(arrowstyle="->", color=RED, lw=1.5))
    right.text(5.0, 0.72, "安全距离是硬边界，无法满足时直接否决原指令", ha="center", color=RED, fontsize=11.5, fontproperties=font_prop, weight="bold")

    fig.text(0.5, 0.035, "原则：无风险时不扰动正常拦截；有风险时只作满足安全约束所需的最小修正。", ha="center", fontsize=12, color=INK, fontproperties=font_prop)
    fig.subplots_adjust(left=0.035, right=0.965, bottom=0.09, top=0.90, wspace=0.045)
    fig.savefig(FIGURES["correction"], dpi=300, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def rounded_box(ax, x: float, y: float, w: float, h: float, text: str, face: str, edge: str, font_prop: font_manager.FontProperties, fontsize: float = 12.0, linewidth: float = 1.6, weight: str = "normal") -> None:
    patch = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.012,rounding_size=0.018", facecolor=face, edgecolor=edge, linewidth=linewidth)
    ax.add_patch(patch)
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fontsize, color=INK, fontproperties=font_prop, weight=weight, linespacing=1.35)


def connect(ax, start: tuple[float, float], end: tuple[float, float], color: str = NAVY, label: str | None = None, label_offset: tuple[float, float] = (0.0, 0.0), font_prop: font_manager.FontProperties | None = None, style: str = "-") -> None:
    arrow = FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=14, linewidth=1.55, color=color, linestyle=style, connectionstyle="arc3,rad=0.0")
    ax.add_patch(arrow)
    if label:
        x = (start[0] + end[0]) / 2 + label_offset[0]
        y = (start[1] + end[1]) / 2 + label_offset[1]
        ax.text(x, y, label, ha="center", va="center", fontsize=10.0, color=color, fontproperties=font_prop, bbox=dict(facecolor="white", edgecolor="none", pad=0.8))


def draw_control_flow(font_prop: font_manager.FontProperties) -> None:
    fig, ax = plt.subplots(figsize=(12.0, 8.0), dpi=300)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.5, 0.955, "有限视场约束下拦截制导与机间安全一体化控制流程", ha="center", va="center", fontsize=20, color=INK, fontproperties=font_prop, weight="bold")

    rounded_box(ax, 0.035, 0.78, 0.18, 0.105, "拍摄时刻同步\n图像、位姿与局部航迹", "#EAF0F6", BLUE, font_prop, 11.7, weight="bold")
    rounded_box(ax, 0.265, 0.78, 0.18, 0.105, "形成相对状态\n视线、距离、速度与不确定度", "#E9F4F1", TEAL, font_prop, 11.7, weight="bold")
    rounded_box(ax, 0.495, 0.78, 0.18, 0.105, "状态机与进入条件\n顺序、方向、高度和视觉质量", "#FFF4DF", ORANGE, font_prop, 11.7, weight="bold")
    rounded_box(ax, 0.725, 0.78, 0.235, 0.105, "正常拦截指令\n比例导引或模型预测控制", "#F0ECF7", PURPLE, font_prop, 11.7, weight="bold")
    connect(ax, (0.215, 0.832), (0.265, 0.832), font_prop=font_prop)
    connect(ax, (0.445, 0.832), (0.495, 0.832), font_prop=font_prop)
    connect(ax, (0.675, 0.832), (0.725, 0.832), font_prop=font_prop)

    rounded_box(ax, 0.70, 0.59, 0.26, 0.105, "短期航迹预测\n交叉点、最近距离和冲突时间", "#EDF2F4", NAVY, font_prop, 11.7, weight="bold")
    connect(ax, (0.842, 0.78), (0.835, 0.695), font_prop=font_prop)

    diamond = Polygon([[0.58, 0.642], [0.655, 0.705], [0.73, 0.642], [0.655, 0.579]], closed=True, facecolor="#FDECEC", edgecolor=RED, linewidth=1.8)
    ax.add_patch(diamond)
    ax.text(0.655, 0.642, "存在\n安全风险？", ha="center", va="center", fontsize=11.2, color=RED, fontproperties=font_prop, weight="bold")
    connect(ax, (0.70, 0.642), (0.73, 0.642), font_prop=font_prop)

    rounded_box(ax, 0.34, 0.56, 0.205, 0.165, "风险触发的安全修正层\n\n控制屏障函数形成硬约束\n二次规划最小改动原指令", "#FFF0F0", RED, font_prop, 11.6, linewidth=2.2, weight="bold")
    connect(ax, (0.58, 0.642), (0.545, 0.642), color=RED, label="是", label_offset=(0.0, 0.025), font_prop=font_prop)

    rounded_box(ax, 0.055, 0.56, 0.215, 0.165, "有限视场与平台边界\n\n目标留在安全像面范围\n检查速度、加速度和高度", "#FFF7E9", ORANGE, font_prop, 11.4, linewidth=1.8, weight="bold")
    connect(ax, (0.34, 0.642), (0.27, 0.642), color=RED, font_prop=font_prop)
    connect(ax, (0.655, 0.579), (0.655, 0.48), color=TEAL, label="否", label_offset=(0.027, 0.0), font_prop=font_prop)

    rounded_box(ax, 0.56, 0.375, 0.225, 0.105, "输出实际控制指令\n并记录安全层是否介入", "#E8F3EB", GREEN, font_prop, 12.0, linewidth=2.0, weight="bold")

    diamond2 = Polygon([[0.41, 0.428], [0.475, 0.478], [0.54, 0.428], [0.475, 0.378]], closed=True, facecolor="#FDECEC", edgecolor=RED, linewidth=1.8)
    ax.add_patch(diamond2)
    ax.text(0.475, 0.428, "约束\n可满足？", ha="center", va="center", fontsize=10.9, color=RED, fontproperties=font_prop, weight="bold")
    connect(ax, (0.54, 0.428), (0.56, 0.428), color=GREEN, label="是", label_offset=(0.0, 0.025), font_prop=font_prop)
    connect(ax, (0.442, 0.56), (0.475, 0.478), color=RED, font_prop=font_prop)
    connect(ax, (0.27, 0.59), (0.41, 0.428), color=ORANGE, font_prop=font_prop)

    rounded_box(ax, 0.17, 0.33, 0.24, 0.105, "安全层否决原指令\n减速分离、退出末段或等待", "#FCE7E7", RED, font_prop, 11.8, linewidth=2.3, weight="bold")
    connect(ax, (0.41, 0.428), (0.41, 0.405), color=RED, label="否", label_offset=(-0.027, 0.012), font_prop=font_prop)

    diamond3 = Polygon([[0.34, 0.215], [0.415, 0.27], [0.49, 0.215], [0.415, 0.16]], closed=True, facecolor="#FFF4DF", edgecolor=ORANGE, linewidth=1.8)
    ax.add_patch(diamond3)
    ax.text(0.415, 0.215, "视觉目标\n仍然有效？", ha="center", va="center", fontsize=10.8, color="#8A5A00", fontproperties=font_prop, weight="bold")
    connect(ax, (0.655, 0.375), (0.49, 0.235), color=NAVY, font_prop=font_prop)

    rounded_box(ax, 0.555, 0.135, 0.215, 0.125, "保持末端接近\n达到规定半径并连续保持\n随后退出或安全等待", "#E8F3EB", GREEN, font_prop, 11.4, linewidth=1.8, weight="bold")
    connect(ax, (0.49, 0.215), (0.555, 0.205), color=GREEN, label="是", label_offset=(0.0, 0.026), font_prop=font_prop)

    rounded_box(ax, 0.095, 0.135, 0.215, 0.125, "退出末段\n按预测方向重搜；证据不足时\n进入安全等待并申请新指令", "#EDF2F4", NAVY, font_prop, 11.2, linewidth=1.8, weight="bold")
    connect(ax, (0.34, 0.215), (0.31, 0.205), color=ORANGE, label="否", label_offset=(0.0, 0.026), font_prop=font_prop)
    connect(ax, (0.29, 0.33), (0.25, 0.26), color=RED, font_prop=font_prop)

    connect(ax, (0.205, 0.135), (0.12, 0.095), color=MUTED, style="--", font_prop=font_prop)
    connect(ax, (0.12, 0.095), (0.12, 0.78), color=MUTED, style="--", label="下一拍重新估计", label_offset=(0.055, 0.0), font_prop=font_prop)
    connect(ax, (0.665, 0.135), (0.84, 0.095), color=MUTED, style="--", font_prop=font_prop)
    connect(ax, (0.84, 0.095), (0.84, 0.59), color=MUTED, style="--", label="闭环反馈", label_offset=(0.042, 0.0), font_prop=font_prop)

    ax.add_patch(FancyBboxPatch((0.79, 0.015), 0.18, 0.055, boxstyle="round,pad=0.01,rounding_size=0.01", facecolor="#FDECEC", edgecolor=RED, linewidth=1.6))
    ax.text(0.88, 0.043, "安全层拥有最终否决权", ha="center", va="center", fontsize=10.8, color=RED, fontproperties=font_prop, weight="bold")
    fig.subplots_adjust(left=0.01, right=0.99, bottom=0.01, top=0.99)
    fig.savefig(FIGURES["flow"], dpi=300, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def draw_lab_route(font_prop: font_manager.FontProperties) -> None:
    fig, ax = plt.subplots(figsize=(12.0, 7.7), dpi=300)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.5, 0.955, "封闭实验室逐级实施与放行路线", ha="center", va="center", fontsize=20, color=INK, fontproperties=font_prop, weight="bold")
    ax.text(0.5, 0.905, "完成判据为进入规定交会半径并连续保持若干帧，全过程禁止真实碰撞", ha="center", va="center", fontsize=12.5, color=RED, fontproperties=font_prop, weight="bold")

    ax.add_patch(FancyBboxPatch((0.035, 0.745), 0.93, 0.105, boxstyle="round,pad=0.01,rounding_size=0.018", facecolor="#EDF2F4", edgecolor=NAVY, linewidth=1.8))
    preparation = "共同安全条件：定位基准与时钟校准　地理围栏　低速限高　独立急停　软质目标　全过程日志与回放"
    ax.text(0.5, 0.797, preparation, ha="center", va="center", fontsize=11.6, color=INK, fontproperties=font_prop, weight="bold")

    stages = [
        (0.04, BLUE, "第一阶段\n2对2", "单一交叉点\n验证基础闭环和安全介入", "零安全距离违规"),
        (0.23, TEAL, "第二阶段\n4对4", "多交叉点与目标转弯\n加入进入顺序和方向协调", "联合指标全部达标"),
        (0.42, ORANGE, "第三阶段\n6对6", "拥挤空域与短时丢失\n加入高度协调和安全等待", "重复试验保持非退化"),
        (0.61, PURPLE, "第四阶段\n进入仿真环境", "扩大速度、时延和丢帧范围\n形成随机场景压力矩阵", "参数边界可复现"),
        (0.80, GREEN, "第五阶段\n更大规模", "按输入规模扩展\n验证计算时延和安全裕度", "满足放行门槛再扩展"),
    ]
    for x, color, title, content, gate in stages:
        w = 0.16
        ax.add_patch(FancyBboxPatch((x, 0.42), w, 0.265, boxstyle="round,pad=0.012,rounding_size=0.018", facecolor="white", edgecolor=color, linewidth=2.1))
        ax.add_patch(Rectangle((x, 0.59), w, 0.095, facecolor=color, edgecolor=color))
        ax.text(x + w / 2, 0.638, title, ha="center", va="center", fontsize=12.2, color="white", fontproperties=font_prop, weight="bold", linespacing=1.25)
        ax.text(x + w / 2, 0.515, content, ha="center", va="center", fontsize=10.5, color=INK, fontproperties=font_prop, linespacing=1.45)
        ax.add_patch(FancyBboxPatch((x + 0.012, 0.438), w - 0.024, 0.047, boxstyle="round,pad=0.006,rounding_size=0.01", facecolor=LIGHT, edgecolor=MID, linewidth=0.9))
        ax.text(x + w / 2, 0.461, gate, ha="center", va="center", fontsize=9.3, color=color, fontproperties=font_prop, weight="bold")

    for i in range(4):
        x1 = stages[i][0] + 0.155
        x2 = stages[i + 1][0]
        connect(ax, (x1 + 0.006, 0.555), (x2 - 0.006, 0.555), color=NAVY, font_prop=font_prop)

    ax.add_patch(FancyBboxPatch((0.055, 0.255), 0.89, 0.105, boxstyle="round,pad=0.012,rounding_size=0.018", facecolor="#FFF7E9", edgecolor=ORANGE, linewidth=1.8))
    ax.text(0.5, 0.307, "每一规模均覆盖：直交、斜交、同向汇聚、目标急转、短时视觉丢失；每轮固定参数并保留原始记录", ha="center", va="center", fontsize=11.4, color=INK, fontproperties=font_prop, weight="bold")

    ax.add_patch(FancyBboxPatch((0.055, 0.095), 0.89, 0.10, boxstyle="round,pad=0.012,rounding_size=0.018", facecolor="#FDECEC", edgecolor=RED, linewidth=2.2))
    ax.text(0.5, 0.145, "任一阶段出现安全距离违规、约束无解或急停异常：立即停止，回放定位原因，不进入下一规模", ha="center", va="center", fontsize=11.7, color=RED, fontproperties=font_prop, weight="bold")
    ax.text(0.5, 0.035, "完成只表示满足预先规定的交会与安全判据，不等于发生物理接触。", ha="center", va="center", fontsize=11.2, color=MUTED, fontproperties=font_prop)
    fig.subplots_adjust(left=0.01, right=0.99, bottom=0.01, top=0.99)
    fig.savefig(FIGURES["route"], dpi=300, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def set_run_font(run, east_asia: str = BODY_FONT, ascii_font: str = ASCII_FONT, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def set_style_font(style, east_asia: str, ascii_font: str, size: float, bold: bool | None = None) -> None:
    style.font.name = ascii_font
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.replace("#", ""))


def set_cell_margins(cell, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "AAB5BC", size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def set_page_number_start(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(start))


def set_section_layout(section, cover: bool = False) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(23.5)
    section.right_margin = Mm(21.5)
    section.header_distance = Mm(9)
    section.footer_distance = Mm(8)
    section.different_first_page_header_footer = cover


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, BODY_FONT, ASCII_FONT, 12.0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.84)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True

    h1 = styles["Heading 1"]
    set_style_font(h1, HEADING_FONT, ASCII_FONT, 16.5, True)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.keep_together = True
    h1.paragraph_format.outline_level = 0

    h2 = styles["Heading 2"]
    set_style_font(h2, HEADING_FONT, ASCII_FONT, 14.0, True)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.keep_together = True
    h2.paragraph_format.outline_level = 1

    if "Equation CN" not in styles:
        equation = styles.add_style("Equation CN", WD_STYLE_TYPE.PARAGRAPH)
    else:
        equation = styles["Equation CN"]
    set_style_font(equation, MATH_FONT, MATH_FONT, 11.5)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.first_line_indent = Cm(0)
    equation.paragraph_format.space_before = Pt(5)
    equation.paragraph_format.space_after = Pt(5)
    equation.paragraph_format.keep_together = True

    caption = styles["Caption"]
    set_style_font(caption, HEADING_FONT, ASCII_FONT, 9.5, True)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_together = True

    if "Table Text CN" not in styles:
        table_style = styles.add_style("Table Text CN", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_style = styles["Table Text CN"]
    set_style_font(table_style, BODY_FONT, ASCII_FONT, 10.5)
    table_style.paragraph_format.first_line_indent = Cm(0)
    table_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table_style.paragraph_format.line_spacing = 1.2
    table_style.paragraph_format.space_after = Pt(0)

    if "Table Header CN" not in styles:
        table_header_style = styles.add_style("Table Header CN", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_header_style = styles["Table Header CN"]
    set_style_font(table_header_style, HEADING_FONT, ASCII_FONT, 10.5, True)
    table_header_style.paragraph_format.first_line_indent = Cm(0)
    table_header_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_header_style.paragraph_format.line_spacing = 1.15
    table_header_style.paragraph_format.space_after = Pt(0)


def add_body(doc: Document, text: str, lead: str | None = None) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    if lead and text.startswith(lead):
        first = paragraph.add_run(lead)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(lead) :])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)


def add_equation(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Equation CN")
    run = paragraph.add_run(text)
    set_run_font(run, east_asia=MATH_FONT, ascii_font=MATH_FONT, size=11.5)


def add_figure(doc: Document, path: Path, caption: str, width_cm: float = 16.0) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph(style="Caption")
    cap_run = cap.add_run(caption)
    set_run_font(cap_run, east_asia=HEADING_FONT, ascii_font=ASCII_FONT, size=9.5, bold=True)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Cm(widths_cm[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.style = doc.styles["Table Header CN"]
        run = paragraph.add_run(header)
        set_run_font(run, east_asia=HEADING_FONT, ascii_font=ASCII_FONT, size=10.5, bold=True, color="FFFFFF")
    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if row_index % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, "F3F6F7")
        for index, (cell, value) in enumerate(zip(row.cells, row_data)):
            cell.width = Cm(widths_cm[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles["Table Text CN"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            set_run_font(run, size=10.5)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.line_spacing = Pt(1)


def add_heading(doc: Document, text: str, level: int, page_break: bool = False) -> None:
    if page_break:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run().add_break(WD_BREAK.PAGE)
    heading = doc.add_paragraph(style=f"Heading {level}")
    run = heading.add_run(text)
    set_run_font(run, east_asia=HEADING_FONT, ascii_font=ASCII_FONT, size=16.5 if level == 1 else 14.0, bold=True)


def add_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(SHORT_TITLE)
    set_run_font(run, east_asia=HEADING_FONT, ascii_font=ASCII_FONT, size=9.0, color="6B7378")
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "B7C0C5")
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run1 = paragraph.add_run("第 ")
    set_run_font(run1, size=9.0, color="6B7378")
    run2 = paragraph.add_run()
    set_run_font(run2, size=9.0, color="6B7378")
    add_field(run2, "PAGE")
    run3 = paragraph.add_run(" 页")
    set_run_font(run3, size=9.0, color="6B7378")


def build_document() -> None:
    doc = Document()
    configure_styles(doc)
    set_section_layout(doc.sections[0], cover=True)

    props = doc.core_properties
    props.title = TITLE
    props.subject = "多机末端接近制导、防碰撞与有限视场统一控制科学问题论证"
    props.author = "MSM项目组"
    props.keywords = "多机拦截；末端制导；有限视场；防碰撞；安全控制"
    props.comments = "科研仿真与技术论证材料"

    for _ in range(3):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run(" ")
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.line_spacing = 1.18
    title_run = title.add_run("多机末端接近制导与防碰撞\n统一科学问题报告")
    set_run_font(title_run, east_asia=HEADING_FONT, ascii_font=ASCII_FONT, size=25.0, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    subtitle_run = subtitle.add_run("任务难点、算法路线与分阶段验证")
    set_run_font(subtitle_run, east_asia=HEADING_FONT, ascii_font=ASCII_FONT, size=15.0, color="4F5B62")

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line.paragraph_format.space_before = Pt(2)
    line.paragraph_format.space_after = Pt(80)
    line_run = line.add_run("━━━━━━━━━━━━━━━━━━━━")
    set_run_font(line_run, east_asia=HEADING_FONT, ascii_font=ASCII_FONT, size=12.0, color=TEAL)

    group = doc.add_paragraph()
    group.alignment = WD_ALIGN_PARAGRAPH.CENTER
    group.paragraph_format.space_after = Pt(8)
    group_run = group.add_run("MSM 项目组")
    set_run_font(group_run, east_asia=HEADING_FONT, ascii_font=ASCII_FONT, size=12.0, bold=True, color=INK)

    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_after = Pt(10)
    date_run = date.add_run("2026 年 8 月")
    set_run_font(date_run, size=11.0, color=MUTED)

    classification = doc.add_paragraph()
    classification.alignment = WD_ALIGN_PARAGRAPH.CENTER
    class_run = classification.add_run("科研仿真与技术论证材料")
    set_run_font(class_run, size=9.5, color="777F84")

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_layout(body_section)
    set_page_number_start(body_section, 1)
    add_header_footer(body_section)

    add_heading(doc, "一、任务、主要难点与关键技术", 1)
    add_heading(doc, "1.1 任务边界与核心矛盾", 2)
    add_body(
        doc,
        "多架拦截无人机在末端空域分别接近已分配目标时，必须同步满足三项条件：在规定时间内进入交会半径，保持目标处于机载相机有效视场，并使任意机对始终具有足够安全间隔。单机比例导引根据本机与目标的相对运动压缩脱靶量，各机独立计算所得航向均可能合理；当多个目标航路交叉、各拦截机提前点相近或前后机速度变化不一致时，多条正常航迹会在有限空间和相近时刻汇聚，形成对穿、追尾或上下穿越风险。末端闭合速度较高，等到当前距离触及门限再避让，剩余控制余量通常已经不足。",
    )
    add_body(
        doc,
        "机间避让与目标观测存在直接耦合。侧向分离、抬升和减速能够增加机间距离，也会改变相机视轴、目标像面位置和闭合速度；持续追求像面居中又会压缩可选避让方向。控制器需要在同一拍摄时刻状态上同时计算目标接近趋势、友机短期位置、视场边缘余量和平台可用控制量，并按照机间安全、地理围栏与平台边界、视场保持、接近效率的固定优先级形成实际指令。安全约束与视场约束发生冲突时，先执行分离和退出动作，视觉重获在安全状态下继续进行。",
    )
    add_body(
        doc,
        "现有研究基础已覆盖比例导引、视觉视线处理、末端状态判定和限定时间的丢帧处置，可生成单机正常接近指令，并为拍摄时刻对齐、相对状态估计和状态切换提供实现依据。控制屏障函数安全层、全体拦截机联合二次规划、进入顺序与方向高度协调以及多机封闭实验室试验尚属拟实施与待验证内容。现有单机或多对独立导引结果只能说明正常制导链具备基础条件，不能替代成员间安全约束和联合完成指标的验证。",
    )

    add_heading(doc, "1.2 主要难点", 2)
    add_body(
        doc,
        "第一类难点来自目标机动与异步观测。目标急转、加减速和高度变化会使匀速外推迅速失准，检测框靠近图像边缘时还会出现裁切和质心偏移。图像对应拍摄时刻，消息到达时刻只反映传输过程；使用到达时刻的机体位姿解释旧图像，会使空间视线、相对速度和预测交叉点产生系统偏差。拍摄时刻、到达时刻、机体位姿、相机安装关系、局部航迹质量和协方差需要完整保留，短时无测量只能在预先设定的外推窗口内使用。",
    )
    add_body(
        doc,
        "第二类难点来自多航迹时空汇聚。几何线段相交并不必然构成即时危险，关键量包括两机到达交叉区域的时间差、预测窗口内最近接近时刻、预计最近距离和不确定度膨胀后的安全裕度。当前间隔较大但闭合速度很高的机对需要提前干预；空间路径接近但到达时间相差较大的机对无需过早改变制导。风险预测还要覆盖通信、求解、执行器响应和完成有效分离所需的总时延。",
    )
    add_body(
        doc,
        "第三类难点来自多重控制约束竞争。正常制导要求压缩脱靶量，视场保持要求限制像面偏差，机间安全要求扩大分离，平台边界又限制速度、加速度、倾角和爬升率。一架无人机的修正动作会改变相邻机下一周期的预测结果，多个安全约束可能同时作用于同一控制量。求解器既要在风险触发后及时改变指令，也要在风险解除后平稳恢复正常制导，并通过确定性的左右、上下和先后规则抑制同向避让、反复切换及控制抖动。",
    )
    add_body(
        doc,
        "第四类难点来自规模增长和约束可解性。两机只产生一个待检查机对，六机产生十五个机对，继续扩展时检查数量近似按机数平方增长。定位误差、时钟误差、通信延迟和执行器迟滞会共同侵蚀标称间隔，安全距离需计入机体包络、误差上界、响应时延和制动能力。联合二次规划必须在一个控制周期内结束；输入失真、计算超时或约束无可行解时，不得沿用未经安全确认的旧指令。",
    )
    add_table(
        doc,
        ["主要矛盾", "直接表现", "需要统一处理的量"],
        [
            ["目标机动与视觉丢失", "视线突变、检测框裁切、短时无测量", "拍摄时刻、相对状态、不确定度和重获时限"],
            ["多航迹汇聚", "交叉、追尾、上下穿越和预计最近距离不足", "交叉点、最近距离、冲突时间和安全裕度"],
            ["制导与避让冲突", "绕开友机后脱靶量增大或目标出框", "正常指令、最小安全修正和视场余量"],
            ["平台与实时边界", "指令饱和、响应滞后、规模增加后计算超时", "速度、加速度、高度、时延和可解性"],
        ],
        [3.5, 5.8, 6.7],
    )
    add_figure(doc, FIGURES["crossing"], "图 1  多航迹交叉三维场景：单机合理航迹在组合后可能形成汇聚风险", 15.8)

    add_heading(doc, "1.3 关键技术", 2)
    add_body(
        doc,
        "针对上述问题，拟重点突破“有限视场约束下拦截制导与机间安全的一体化控制技术”。正常制导层依据目标相对运动生成比例导引指令，平台模型和预测条件成熟后扩展模型预测控制；联合预测层对全部相关机对计算交叉位置、最近接近时刻、预计最近距离及其不确定度；安全修正层以控制屏障函数建立机间安全硬约束，通过联合二次规划求取偏离正常指令最小的实际控制量；状态机负责进入、跟踪、风险修正、重搜、安全等待、完成和退出之间的条件切换。四部分共享拍摄时刻状态、控制边界和日志定义。",
    )
    add_body(
        doc,
        "实际控制指令须同时满足机间安全、平台可飞和当前状态有效三项边界。预测无风险时，安全修正量保持为零或数值容差内的小量；风险触发后，修正仅覆盖恢复安全所需的方向、速度或高度变化。安全层拥有最终否决权，正常制导、视场保持和安全距离无法同时满足时，立即转入减速分离、确定性方向或高度分离、退出末段和安全等待。该否决结果同步进入状态转移、控制输出、事件日志与试验判定，任何上级进入计划均不得绕过实时安全检查。",
    )

    add_heading(doc, "二、具体算法步骤与实施流程", 1, page_break=True)
    add_heading(doc, "2.1 拍摄时刻同步、相对状态与正常制导", 2)
    add_body(
        doc,
        "每个控制周期以图像拍摄时刻整理输入。机体位置、速度、姿态和相机安装关系插值到该时刻，检测框中心经相机内参和姿态变换转换为空间单位视线；局部航迹提供视线变化率、尺度变化、连续观测长度和质量分数，中心航迹或多帧估计提供距离、目标速度及协方差。只有方位信息时，系统保留较大的距离不确定度，假定距离不得进入确定性安全判断。相对状态定义为：",
    )
    add_equation(doc, "r_i(t_m)=p_T(t_m)-p_I,i(t_m)，  v_rel,i=v_T(t_m)-v_I,i(t_m)，  λ_i=r_i/||r_i||")
    add_body(
        doc,
        "式中，t_m表示图像拍摄时刻；p_T与p_I,i分别表示目标和第i架拦截机在统一坐标系中的位置；v_T与v_I,i分别表示对应速度；r_i、v_rel,i和λ_i分别表示相对位置、相对速度和目标单位视线。其作用是把图像、位姿和航迹归并到同一物理时刻，为闭合速度、视线角速率、视场余量及短期位置预测提供一致输入。信息龄期、连续丢帧和目标机动程度进入协方差传播，预测边界随不确定度增大而扩展。",
    )
    add_body(
        doc,
        "正常拦截层以现有比例导引为基础，根据视线转动和闭合速度产生法向加速度，使拦截机对目标运动趋势提前修正。三维比例导引指令为：",
    )
    add_equation(doc, "a_i^0=N_i V_c,i(ω_i×λ_i)，  V_c,i=-v_rel,i^T λ_i")
    add_body(
        doc,
        "式中，a_i^0表示第i架拦截机未经机间安全检查的正常加速度指令；N_i表示导引系数；V_c,i表示闭合速度；ω_i表示视线角速度；λ_i表示目标单位视线。其作用是把视线转动转换为侧向和垂向修正，闭合速度越大、视线转动越快，正常指令越早建立拦截提前量。模型预测控制作为拟扩展路线，在有限预测域内联合考虑距离误差、像面余量、指令平滑和平台边界，所得首个控制量仍作为正常指令送入外层安全检查。",
    )

    add_heading(doc, "2.2 短期冲突预测与安全修正", 2)
    add_body(
        doc,
        "正常指令形成后，对相关空域内全部拦截机进行短期联合预测。预测窗口覆盖观测传输、状态计算、控制求解、执行器响应和完成一次有效分离所需时间，并按平台机动能力设置上限。对任意机对i、j，以相对位置p_ij和相对速度v_ij估计窗口内最近接近时刻与预计最近距离：",
    )
    add_equation(doc, "t_cpa=clip[-(p_ij^T v_ij)/||v_ij||²，0，T_p]，  d_cpa=||p_ij+v_ij t_cpa||")
    add_body(
        doc,
        "式中，T_p表示短期预测窗口；clip表示把计算时刻限制在0至T_p范围内；t_cpa表示最近接近时刻；d_cpa表示该时刻的预计机间距离。其作用是区分仅有空间交叉的航迹与在相近时刻进入同一区域的真实冲突。风险判定同时计算预测交叉点、到达时间差和协方差膨胀后的距离下界；距离下界低于预先确定的安全距离且冲突时间进入响应窗口时，安全层提前介入。",
    )
    add_body(
        doc,
        "风险触发后，拟实施的控制屏障函数安全层把规定安全距离转换为不可放松的控制约束。对机对i、j定义安全函数h_ij，并对加速度控制系统设置相对阶次为二的屏障条件：",
    )
    add_equation(doc, "h_ij=||p_i-p_j||²-d_safe²，  ḧ_ij+k_1 ḣ_ij+k_0 h_ij≥0")
    add_body(
        doc,
        "式中，p_i和p_j表示两机位置；d_safe表示阶段开始前确定的安全距离；h_ij表示安全裕度函数；k_0和k_1表示正的屏障增益；ḣ_ij和ḧ_ij分别表示安全裕度的一阶、二阶变化率。其作用是同时约束当前距离和接近趋势，使快速闭合的机对在越界前获得足够修正时间，并在分离趋势稳定后逐步释放控制权。d_safe由机体包络、定位误差、时钟与通信误差、控制迟滞及最大制动距离共同确定。",
    )
    add_body(
        doc,
        "联合二次规划在全部活动机对的屏障约束下，求取偏离正常指令最小的实际控制量。拟实施的目标函数与约束为：",
    )
    add_equation(doc, "min_a Σ_i ||a_i-a_i^0||²_W+ρ||s_视场||²；  满足：安全屏障硬约束、视场与动力学边界")
    add_body(
        doc,
        "式中，a表示全体拦截机的实际控制量；a_i^0表示第i架拦截机的正常指令；W表示控制修正权重；s_视场表示视场软约束松弛量；ρ表示视场偏离惩罚系数。其作用是在安全屏障硬约束和平台动力学边界内保持正常拦截意图，仅对风险相关控制分量进行必要修正。视场松弛量只影响目标接近像面边缘的代价，不得削弱机间安全约束。无可行解、计算超时或输入可信度不足时，安全层最终否决正常指令，立即执行减速分离、方向或高度分离、退出末段和安全等待。",
    )
    add_figure(doc, FIGURES["correction"], "图 2  正常制导与安全修正原理：仅在风险触发时最小改动原指令", 16.0)

    add_heading(doc, "2.3 进入顺序、方向与高度协调", 2)
    add_body(
        doc,
        "拥挤末端仅依靠逐周期避让容易造成反复修正，需要在进入前增加时空协调。各机预计到达时间、目标威胁程度、视场余量、剩余机动能力和当前安全裕度共同决定进入顺序，裕度不足的成员在外圈减速或等待；相邻航迹分配不同接近方位，降低同向重叠和对穿概率；场地与平台允许时设置临时高度层或垂向偏置，为平面内难以分开的航迹增加净空。顺序、方向和高度只限定计划进入边界，实时安全层持续审核实际状态并保留否决权。",
    )
    add_body(
        doc,
        "控制优先级依次为机间安全和独立急停、地理围栏与平台可飞边界、有限视场、接近效率和指令平滑。低优先级目标在约束冲突时让位于高优先级边界。左右及上下避让依据稳定机号顺序、相对方位和预设规则确定，进入门限与退出门限保留迟滞区，避免多机同时选择同一方向或在阈值附近往复切换。连续多个机对同时触发时，联合求解一次处理全部活动约束；安全等待为拥挤状态提供可恢复出口。",
    )
    table_page_break = doc.add_paragraph()
    table_page_break.paragraph_format.space_after = Pt(0)
    table_page_break.add_run().add_break(WD_BREAK.PAGE)
    add_table(
        doc,
        ["优先级", "控制要求", "冲突时的处理"],
        [
            ["一", "机间安全距离和独立急停", "具有最终否决权；无可行解时退出末段"],
            ["二", "地理围栏、速度、加速度和高度边界", "限制可选控制范围，不允许以越界换取接近"],
            ["三", "目标位于有限视场安全区", "优先减速、转向保视场；与安全冲突时允许退出"],
            ["四", "接近效率、时间误差和指令平滑", "在前三项满足后作优化"],
        ],
        [2.1, 6.0, 7.9],
    )

    add_heading(doc, "2.4 有限视场、视觉丢失与状态机", 2)
    add_body(
        doc,
        "有限视场约束在像面直接计算。目标预测像点相对图像中心的水平、垂直偏差分别记为e_u和e_v，图像半宽与半高分别记为U和V，边缘安全量记为m_f，对应约束为：",
    )
    add_equation(doc, "|e_u|≤U-m_f，  |e_v|≤V-m_f")
    add_body(
        doc,
        "式中，e_u和e_v表示目标预测像点的水平与垂直偏差；U和V表示可用像面边界；m_f表示为目标机动、姿态误差和控制迟滞预留的视场余量。其作用是阻止目标仅以贴近图像边缘的状态进入末段，并为下一周期观测保留可控空间。比例导引链通过速度调度、偏航和高度修正保持余量，模型预测控制可在多步预测中直接约束未来像点。安全修正必然造成目标出框时，机间安全保持最高优先级，状态机随即退出末段。",
    )
    add_body(
        doc,
        "视觉短时丢失期间，只在限定帧数和限定时间内使用最近可靠视线及其变化趋势进行减速外推，同时继续预测全部机对距离。超过外推窗口、图像质量持续低于门限或目标关系不稳定时，状态机退出末段；预测方向可信时按限定扇区重搜，方向证据不足时保持安全航向、速度和高度，等待新观测与新进入指令。目标重获后重新检查连续观测长度、视场余量、闭合速度、任务版本和联合约束可解性，全部条件满足后方可恢复末端接近。",
    )
    add_body(
        doc,
        "状态机包括等待、正常跟踪、末段接近、风险修正、视觉重搜、安全等待、完成和退出八类状态。每次转移记录触发原因、拍摄时刻、到达时刻、预测最近距离、实际最小距离、正常指令、修正指令和安全层介入标志。完成状态仅在进入规定交会半径并连续保持规定帧数后成立，安全距离违规、急停触发、输入失真和视觉证据不足均排除完成判定。控制流程按观测更新、正常制导、冲突预测、安全修正、实际执行和下一拍复核循环运行。",
    )
    doc.paragraphs[-1].paragraph_format.keep_together = True
    add_figure(doc, FIGURES["flow"], "图 3  完整控制流程：安全层位于正常制导与实际指令之间并具有最终否决权", 16.0)

    add_heading(doc, "三、期望结果与后续实施", 1, page_break=True)
    add_heading(doc, "3.1 可量测指标与判定条件", 2)
    add_body(
        doc,
        "技术验证以可记录、可复核计算的联合指标为依据。交会半径、连续保持帧数、安全距离、视场边缘余量、求解期限和退出时限在每一阶段开始前确定，试验过程中不得根据结果临时调整。表中百分比和时延比例为首轮建议目标，后续可依据平台尺寸、定位精度、场地尺度和控制周期修订；修订前后的配置、原始记录和原因说明均需保留。安全距离与交会半径分别服务于机间安全和任务完成，两项门限独立设置。",
    )
    add_table(
        doc,
        ["指标", "计算方法", "建议判定口径"],
        [
            ["交会完成率", "进入规定交会半径 R_c 并连续保持 K 帧的案例数占比", "首轮建议 K=10 帧；不要求也不允许真实碰撞"],
            ["最小机间距离", "全程所有机对距离的最小值", "不得小于阶段开始前确定的 d_safe，违规次数为零"],
            ["安全裕度", "最小机间距离减去 d_safe", "全程不为负，并报告最小值和低分位值"],
            ["视场保持率", "末段有效目标像点帧数除以末段总帧数", "首轮建议不低于95%，同时报告最长连续丢失"],
            ["重获或安全退出时延", "从视觉失效到重获或进入安全等待的时间", "报告中位值、95分位值和最大值"],
            ["安全修正量", "实际指令与原始指令差的范数及相对比例", "无风险帧应接近零，风险帧报告峰值与持续时间"],
            ["求解时延与无解率", "二次规划逐周期耗时及无可行解次数", "95分位不超过控制周期的20%，不得跨周期沿用危险旧指令"],
            ["联合完成率", "同时满足交会完成、零安全违规和视场口径的案例占比", "按2对2、4对4、6对6分别统计，不混合分母"],
        ],
        [3.1, 6.5, 6.4],
    )
    add_body(
        doc,
        "每个案例保存统一时钟下的图像拍摄时刻、消息到达时刻、原始观测、相对状态、协方差、正常指令、风险判定、修正指令、状态转移和机对距离曲线。评价结果同时给出目标侧完成情况和机对侧安全情况，任何单机进入交会半径均不能替代全体相关机对的安全说明。安全层未介入的案例需保留无风险预测证据，介入案例需给出修正前后预计最近距离、实际最小距离、控制改变量及接近效果变化。",
    )

    add_heading(doc, "3.2 分阶段验证与封闭实验室路线", 2)
    add_body(
        doc,
        "封闭实验室试验全过程禁止真实碰撞。目标使用软质标志、虚拟目标点或受保护牵引目标，拦截机进入规定交会半径并连续保持若干帧后，按状态机退出或进入安全等待。场地配置定位基准与时钟校准、地理围栏、低速限高、独立急停、人工监护和全过程回放。机间距离越界、定位质量异常、联合约束无解或急停链路异常均立即终止当轮，机体接触不作为完成证据。",
    )
    add_body(
        doc,
        "验证按2对2、4对4、6对6逐级推进，每一规模覆盖直交、斜交、同向汇聚、目标急转和单机短时视觉丢失五类工况，每类建议不少于10组独立初始条件。2对2检查屏障约束方向、最小安全修正和最终否决；4对4检查多冲突并发时的进入顺序、方向协调与确定性避让；6对6检查拥挤空域中的高度协调、连续风险触发、无可行解处置和安全等待。联合完成率、安全距离、视场保持与求解时延全部达到预设门槛后，方可进入下一规模。",
    )
    add_body(
        doc,
        "完成6对6封闭实验室基线后，在AirSim中扩大速度、目标机动、相机时延、通信延迟和连续丢帧范围，以固定随机种子形成可重复压力矩阵。仿真继续执行相同的交会半径、连续帧、安全距离和状态机判定，几何距离指标保持为主要安全依据。2对2、4对4和6对6仅为基线场景，算法规模由输入机数决定；扩展试验重点统计机对数量增长引起的求解时延、无解率、安全层介入率和安全裕度变化。",
    )
    add_figure(doc, FIGURES["route"], "图 4  封闭实验室实施路线：逐级放行、失败回退、完成后再进入仿真和更大规模", 15.5)

    add_heading(doc, "3.3 期望形成的工程结果与边界", 2)
    add_body(
        doc,
        "预期形成四类工程结果：形成拍摄时刻一致的相对状态、风险预测和控制指令定义，坐标方向、时间来源与协方差可逐项追溯；形成正常制导与安全修正分层的控制程序，明确无风险保持、风险最小修正和无可行解否决三类行为；形成覆盖进入顺序、方向与高度协调、有限视场、视觉丢失、重搜和安全等待的状态机；形成按规模、工况和随机种子组织的试验记录，可逐案例复算交会完成、安全距离、视场保持和计算时延。",
    )
    add_body(
        doc,
        "当前能力基础限于比例导引、视觉视线处理、末端状态判定和短时丢帧处置。控制屏障函数安全层尚待实现，联合二次规划与多机协调尚待集成，2对2至6对6封闭实验室试验尚未开展，相关指标没有形成验证结果。后续结论须注明实际阶段、场景、样本数量、预先确定的参数和判定条件；AirSim结果属于仿真证据，封闭实验室进入交会半径也不等同于真实对抗或外场实装能力。",
    )
    add_body(
        doc,
        "工程阶段的判定条件为：目标机动、短时视觉丢失和多航迹汇聚同时出现时，系统能够依靠可复核的最小修正保持安全距离；约束无法同时满足时，安全层能够否决正常指令并进入可恢复状态；上述行为在2对2、4对4、6对6、仿真压力矩阵和后续扩展规模中按同类指标重复成立。任一阶段出现安全距离违规、无解处置失效或急停异常，立即停止扩展并回放定位原因。",
    )

    doc.save(OUTPUT_DOCX)


def all_docx_text(doc: Document) -> str:
    chunks = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


def verify_outputs() -> dict[str, int]:
    if not OUTPUT_DOCX.exists():
        raise FileNotFoundError(OUTPUT_DOCX)
    for path in FIGURES.values():
        if not path.exists():
            raise FileNotFoundError(path)
        with Image.open(path) as image:
            width, height = image.size
            if width < 3000 or height < 1700:
                raise AssertionError(f"Figure resolution too small: {path.name} {width}x{height}")
            if image.format != "PNG":
                raise AssertionError(f"Figure is not PNG: {path}")

    doc = Document(OUTPUT_DOCX)
    text = all_docx_text(doc)
    chinese_characters = len(re.findall(r"[\u3400-\u4dbf\u4e00-\u9fff]", text))
    if not 4000 <= chinese_characters <= 6000:
        raise AssertionError(f"Chinese character count outside requested range: {chinese_characters}")
    sensitive_patterns = (
        "\u672c\u62a5\u544a",
        "\u95ee\u9898\u7684\u5b9e\u8d28\u4e0d\u662f",
        "\u6838\u5fc3\u4e0d\u662f",
        "\u901a\u4fd7\u5730\u8bf4",
        "\u7b80\u5355\u6765\u8bf4",
        "\u6362\u8a00\u4e4b",
        "\u5199\u6210",
        "\u7edf\u4e00\u8868\u8ff0\u4e3a",
        "\u8d2f\u901a",
        "\u6536\u53e3",
        "\u8d4b\u80fd",
        "\u6253\u9020",
        "\u5408\u540c",
        "\u69fd\u4f4d",
    )
    sensitive_phrase_count = sum(text.count(pattern) for pattern in sensitive_patterns)
    contrast_template_count = len(re.findall(r"\u4e0d\u662f[^\u3002\uff1b\n]{0,80}\u800c\u662f", text))
    if sensitive_phrase_count or contrast_template_count:
        raise AssertionError(
            "Sensitive wording found in report: "
            f"phrases={sensitive_phrase_count}, contrast_templates={contrast_template_count}"
        )
    chapter_headings = [
        "一、任务、主要难点与关键技术",
        "二、具体算法步骤与实施流程",
        "三、期望结果与后续实施",
    ]
    for heading in chapter_headings:
        if text.count(heading) != 1:
            raise AssertionError(f"Expected chapter heading exactly once: {heading}")
    heading_one_texts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.style.name == "Heading 1"]
    if heading_one_texts != chapter_headings:
        raise AssertionError(f"Unexpected body chapter structure: {heading_one_texts}")
    if len(doc.tables) != 3:
        raise AssertionError(f"Expected exactly three tables, got {len(doc.tables)}")
    if len(doc.inline_shapes) != 4:
        raise AssertionError(f"Expected exactly four embedded figures, got {len(doc.inline_shapes)}")
    equation_count = sum(1 for paragraph in doc.paragraphs if paragraph.style.name == "Equation CN")
    if equation_count != 6:
        raise AssertionError(f"Expected exactly six displayed equations, got {equation_count}")
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.style.name != "Equation CN":
            continue
        explanation = next(
            (candidate.text.strip() for candidate in doc.paragraphs[index + 1 :] if candidate.text.strip()),
            "",
        )
        if not explanation.startswith("式中，") or "其作用是" not in explanation:
            raise AssertionError(f"Equation lacks a formal explanation: {paragraph.text}")
    required_terms = (
        "有限视场约束下拦截制导与机间安全的一体化控制技术",
        "安全层拥有最终否决权",
        "图像拍摄时刻",
        "控制屏障函数",
        "联合二次规划",
        "无可行解",
        "拟实施与待验证",
    )
    missing_terms = [term for term in required_terms if term not in text]
    if missing_terms:
        raise AssertionError(f"Required technical statements missing: {missing_terms}")

    with ZipFile(OUTPUT_DOCX) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise AssertionError(f"Corrupt DOCX member: {bad_member}")
        required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
        if not required.issubset(set(archive.namelist())):
            raise AssertionError("DOCX is missing required package members")
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        if len(media) != 4 or any(not name.lower().endswith(".png") for name in media):
            raise AssertionError(f"Unexpected embedded media: {media}")

    nonempty_paragraphs = sum(1 for paragraph in doc.paragraphs if paragraph.text.strip())
    narrative_paragraphs = sum(
        1
        for paragraph in doc.paragraphs
        if paragraph.text.strip() and paragraph.style.name == "Normal" and not paragraph.text.startswith("图 ")
    )
    return {
        "chinese_characters": chinese_characters,
        "nonempty_paragraphs": nonempty_paragraphs,
        "narrative_paragraphs": narrative_paragraphs,
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "equations": equation_count,
        "sensitive_phrases": sensitive_phrase_count,
        "contrast_templates": contrast_template_count,
        "sections": len(doc.sections),
    }


def main() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    font_prop = configure_matplotlib()
    draw_crossing_scene(font_prop)
    draw_nominal_and_correction(font_prop)
    draw_control_flow(font_prop)
    draw_lab_route(font_prop)
    build_document()
    stats = verify_outputs()
    figure_stats = []
    for path in FIGURES.values():
        with Image.open(path) as image:
            dpi = tuple(round(value) for value in image.info.get("dpi", (0, 0)))
            figure_stats.append(f"{path.name}={image.width}x{image.height}@{dpi[0]}dpi")
    print(f"OUTPUT={OUTPUT_DOCX}")
    print("STATS=" + ",".join(f"{key}:{value}" for key, value in stats.items()))
    print("FIGURES=" + ";".join(figure_stats))


if __name__ == "__main__":
    main()
