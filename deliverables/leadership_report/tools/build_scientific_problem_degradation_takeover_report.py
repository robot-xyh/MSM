#!/usr/bin/env python3
"""Build the Word-only leadership report on communication-limited takeover."""

from __future__ import annotations

import argparse
import re
import warnings
from pathlib import Path
from typing import Iterable
from zipfile import ZipFile

import matplotlib

matplotlib.use("Agg")
warnings.filterwarnings("ignore", message="Unable to import Axes3D.*")

import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import Circle, FancyBboxPatch, Polygon, Rectangle
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[3]
REPORT_ROOT = REPO_ROOT / "deliverables" / "leadership_report"
ASSET_DIR = REPORT_ROOT / "assets" / "scientific_problem_degradation_takeover"
OUTPUT = REPORT_ROOT / "通信受限条件下一致分配与降级接管科学问题报告_CN.docx"

TITLE = "通信受限条件下一致分配与降级接管科学问题报告"
CHAPTER_TITLES = (
    "一、任务、主要难点与关键技术",
    "二、具体算法步骤与实施流程",
    "三、期望结果与后续实施",
)
FIGURES = (
    ("01_communication_interruption_scenario.png", "通信中断条件下的任务保持与直接接管场景"),
    ("02_temporary_center_election_principle.png", "临时中心机固定排序与过半确认原理"),
    ("03_degradation_recovery_full_flow.png", "中心失效、临时接管、再次失联与恢复交还完整流程"),
    ("04_network_fault_injection_route.png", "网络故障注入与室内分级验证实施路线"),
)

BODY_FONT = "宋体"
HEADING_FONT = "黑体"
LATIN_FONT = "Times New Roman"
FIGURE_FONT_PATH = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")

WORD_INK = "202833"
WORD_BLUE = "1F4E78"
WORD_TEAL = "176B73"
WORD_MUTED = "5F6B78"
WORD_LIGHT = "EAF0F4"
WORD_LINE = "CAD3DC"

BLUE = "#245B78"
TEAL = "#287A78"
GREEN = "#3F7E5E"
ORANGE = "#C06A32"
RED = "#A8453C"
INK = "#202833"
MUTED = "#65727E"
LINE = "#C8D1D8"
LIGHT = "#F4F6F7"
PALE_BLUE = "#EAF2F7"
PALE_TEAL = "#E8F3F1"
PALE_GREEN = "#EBF3ED"
PALE_ORANGE = "#FAF0E8"
PALE_RED = "#F8EDEC"


CHAPTER_ONE: tuple[tuple[str, str], ...] = (
    (
        "",
        "多无人机拦截进入协同执行阶段后，中心节点持续汇总航迹、资源和任务进度，并向各机发布分配计划。地形遮挡、电磁干扰、链路拥塞或设备故障会使消息延迟、丢失、乱序，严重时形成持续断链和网络分区。各机此时掌握的航迹时刻、资源状态和计划版本可能不同，若依据局部信息自行换令，容易造成同一目标被重复占用、协同成员彼此不一致或执行完毕的任务再次下发。",
    ),
    (
        "",
        "任务处置需要同时满足执行连续和发布来源唯一两个约束。中心消息短时波动时，频繁切换发布者会放大扰动并中断原任务；中心持续失效时，长期等待又会错过目标变化和资源调整窗口。有效期内、身份明确且继续执行更安全的任务应保持，来源不明、版本倒退、关系冲突或已经过期的任务应拒绝或暂停，新的分配只能在发布权与任务关系均可核验后生效。",
    ),
    (
        "",
        "体系不设置二级节点，也不预先指定长期备份指挥机。中心失效得到持续证据和成员确认后，由拦截无人机直接推举一架临时中心机，限时承担状态汇总、任务调整和计划发布。推举基准固定为中心最后共同确认的成员清单，连通成员达到该清单过半数才具备接管条件；未过半的分区只维持分区前已确认、仍在有效期内且继续执行更安全的任务。图1给出了中心失效、网络分区和任务保持同时出现时的处置关系。",
    ),
    (
        "",
        "中心失效判断需要区分链路瞬时抖动、消息排队和真正的持续失效。单个心跳迟到不足以改变发布权，固定超时过长又会压缩接管后的任务调整时间；各机若按不同本地时钟独立判断，还可能出现部分节点开始推举、另一部分继续接受原中心计划的矛盾。判断依据必须同时包含消息生成时刻、到达时刻、连续缺失窗口和同伴观察，迟到旧包、重复包及版本倒退消息均不能清除失效计时。",
    ),
    (
        "",
        "直接推举的主要风险是两个候选依据不同局部状态同时取得发布权。通信质量、剩余航时和任务负荷随位置变化，各机若按连续评分各自计算，微小差异就可能改变排序；网络分区后若临时缩小成员范围，两侧还可能分别形成局部多数。候选评价必须绑定同一固定成员清单和同一批带时间标记的状态摘要，先执行硬条件筛选，再按所有节点均能独立核对的固定顺序比较，并以过半确认作为生效条件。",
    ),
    (
        "",
        "临时中心机当选只解决发布来源问题，尚需约束计划替换、任务确认和恢复交还。链路恢复后，旧中心计划与临时计划可能乱序到达；临时中心再次失联时，剩余成员还要在更高接管轮次重新推举；原中心恢复时，其掌握的任务进度可能落后于现场。每份计划因此必须携带接管轮次、计划版本、生成时刻和有效期，恢复阶段还需并行核对航迹、资源、计划与执行状态，冲突消除后再交还发布权。",
    ),
    (
        "",
        "针对上述问题，拟重点突破“时变通信条件下唯一任务关系保持与安全恢复技术”。该技术以最后共同确认的固定成员清单为推举基准，以事件触发交换和持续失效判断形成接管证据，以固定候选排序、接管轮次和过半确认限定发布权，再通过递增计划版本、有效期、冲突拒绝和恢复核对保持任务关系连续。任一环节证据不足时，系统维持已确认的安全任务或进入等待，不发布无法核验的新关系。方案中的直接推举和室内验证尚待实施；现有版本、时效、拒绝和故障注入基础可复用。",
    ),
)


CHAPTER_TWO_BEFORE_FIGURE_TWO: tuple[tuple[str, str], ...] = (
    (
        "",
        "算法按中心正常、失效待确认、推举、临时接管、分区保持、恢复核对和交还七类状态运行。各机持续保存当前发布者、固定成员清单、接管轮次、计划版本、生成时刻与有效截止时刻，任何状态变化均与这组基准比较。完整流程依次处理事件触发交换、持续失效判断、候选排序、过半确认、计划发布、分区保持、再次推举和恢复交还，每次改变均记录依据、确认与拒绝结果。",
    ),
    (
        "",
        "事件触发交换传递航迹、资源和计划三类摘要。航迹摘要包含目标编号、位置速度概况、测量时刻、到达时刻、不确定度和跟踪状态；资源摘要包含无人机编号、剩余航时、通信邻接、计算余量、位置、任务负荷和安全限制；计划摘要包含发布者、接管轮次、版本、生成时刻、有效期、任务关系和确认情况。航迹变化越过门限、目标进入关键阶段、航时跨级、任务完成或退出、通信邻接改变时立即发送，无事件时仅低频发送最小保活摘要；接收端按生成时刻和序号排序，过旧、倒退或超期信息只进入审计记录。",
    ),
    (
        "",
        "中心消息状态分为新鲜、可疑和失效待确认三个阶段，前两个阶段均不改变发布权。消息更新时间超过正常间隔后进入可疑状态，各机提高摘要交换频率；新鲜消息持续缺失达到规定窗口，并由固定成员清单中过半成员报告一致观察后，中心失效才成立。持续窗口同时受最少交换次数和最短持续时间约束，只有来源可核验、生成时刻更新且计划关系连续的中心消息可以复位计时，迟到旧包、重复包和版本倒退消息均保留原计时。",
    ),
    (
        "",
        "进入失效待确认阶段后，各机固定采用中心最后共同确认的成员清单，并对清单内容计算一致的校验值。过半门槛始终以该清单的成员总数为分母，推举期间短时失联成员仍保留在清单中，防止分区两侧按各自在线人数降低门槛。确有成员永久退出时，只能在没有活动推举、网络状态持续稳定并取得原清单过半确认后更新成员清单；条件不足时保持原分母，即使暂时无法接管也不产生两个发布来源。",
    ),
    (
        "",
        "候选首先接受硬条件筛选：状态摘要新鲜，能够与过半成员直接或经转发互通，剩余航时覆盖接管时段，计算资源满足汇总与分配需要，位置和姿态不影响原任务安全。通过筛选后，按通信覆盖、剩余航时、计算能力、位置条件和任务负荷依次形成固定候选排序，各项先划分离散等级，再按既定优先级逐项比较，避免测量微差频繁改变结果。同档候选按预先公布的无人机编号确定先后，排序输入、等级边界和消解规则随任务初始化一并下发，推举过程中不临时修改。",
    ),
)


CHAPTER_TWO_AFTER_FIGURE_TWO: tuple[tuple[str, str], ...] = (
    (
        "",
        "中心失效确认成立后，各机把接管轮次递增一轮，并广播固定成员清单校验值、候选排序首位及所依据摘要的最新时刻。各机计算结果一致时，在本轮只对一名候选发出带来源的确认；结果不一致时先交换缺失摘要，待输入满足时效要求后重新计算，最先广播者不因时序领先取得额外资格。候选获得固定成员总数过半的有效确认后，发布接管轮次、确认成员、证据摘要和授权截止时刻，其他节点完成同项核验后才承认临时中心机。",
    ),
    (
        "",
        "同一固定成员清单和同一接管轮次内，每名成员最多确认一名候选，生效门槛取成员总数的一半向下取整后加一。若两名候选都声称取得过半确认，两个确认集合的成员数之和必然大于固定成员总数，因此二者至少有一名共同成员；该成员若同时确认两名候选，就违反同轮单次确认约束，两个候选不能同时形成合法发布权。成员清单校验值保证两组确认使用相同分母，接管轮次隔离前后推举，授权有效期使失去过半联系的临时中心机到期停止发布新计划。",
    ),
    (
        "",
        "临时中心机汇总最后共同计划、各机执行状态、有效航迹和资源状态，把任务分为保持、重算和暂停三类。身份明确、资源可达且继续执行风险较低的关系优先保持；资源失效、目标状态显著变化或计划临近到期的任务进入重算；来源不明、重复申报或执行状态无法核验的任务暂停。重算先排除目标不可达、剩余航时不足、信息过旧和资源已占用的组合，再综合预计到达时间、目标紧迫度、航迹不确定度、通信风险和换令代价形成分配，唯一性、可达性与时效检查通过后进入成员确认。",
    ),
    (
        "",
        "临时计划携带发布者编号、接管轮次、计划版本、生成时刻、有效开始与截止时刻、被替代版本、固定成员清单校验值和任务内容摘要。接收端先比较接管轮次，再核对发布者资格、版本递增、内容摘要和有效期：较低轮次直接拒绝，较高轮次缺少合法过半确认时同样拒绝，同一轮次与版本出现不同内容则标记为冲突。字段缺失、生成时刻过旧或授权已经到期的计划只进入审计记录；必要执行成员确认相同内容后任务关系生效，未确认成员维持原安全状态。",
    ),
    (
        "",
        "网络分区后，取得固定成员清单过半联系的一侧可以进入推举，但只调整状态能够核验的资源和未被有效关系占用的任务。多数侧不得占用少数侧仍在有效期内的任务，少数侧不能缩小成员清单或产生新分配，只能保持分区前已确认、仍在有效期内且继续执行风险较低的任务；计划到期或安全条件不再满足时转入预定保持状态。分区恢复后先交换任务摘要和执行记录，重复占用、完成状态与版本冲突全部消除，并由更高接管轮次确认后才恢复任务调整。",
    ),
    (
        "",
        "候选状态缺失、清单校验值不一致、关键确认丢失或可通信成员不足都会使推举超时。在线成员数量减少时，过半门槛仍按固定成员清单计算，各机保持最后确认且继续执行风险较低的任务，待重算任务进入等待或预定安全机动，任何候选不得提前发布新关系。固定等待时间结束后，本轮推举关闭，接管轮次递增，失去新鲜状态的候选从硬筛选中排除并再次排序；新一轮仍未过半时继续保持并记录具体阻塞原因。",
    ),
    (
        "",
        "临时中心机停止更新后，各机沿用中心失效的持续窗口检查，不以单次丢包立即撤销授权。持续缺失达到门槛并取得固定成员清单过半确认后，旧授权停止续期，未生效计划撤销，已经确认且继续执行风险较低的任务保持至原有效期截止。剩余成员进入更高接管轮次，重新执行硬筛选、固定排序和过半确认；旧临时中心机稍后恢复时只作为普通成员加入当前轮次，若同时发生分区，仍执行多数侧可推举、少数侧只保持的约束。",
    ),
    (
        "",
        "原中心恢复发送后先以观察者身份加入当前状态，不立即覆盖临时中心机。并行核对期内，原中心、临时中心机和各执行成员交换航迹、资源、任务、执行结果、接管轮次、计划版本与有效期，连续达到规定核对窗口且结论一致后才进入交还。原中心随后发起更高接管轮次的交还请求，取得过半确认并从明确时刻发布新版本，临时中心机收到交还确认后停止续期；目标身份、资源占用、完成状态或版本仍有冲突时，当前发布权保持，冲突任务暂停核对。",
    ),
    (
        "",
        "中心失效只启动推举，候选过半只授予限时发布权，临时计划经过版本、有效期和必要成员确认后才进入执行。推举超时、临时中心再次失联、网络分区和恢复冲突均返回保持状态或进入更高接管轮次，不降低固定成员清单下的过半门槛。图3把失效判断、直接推举、计划生效、分区保持、再次推举和恢复交还置于同一状态流程，每个转移都能由消息时刻、确认记录、轮次和版本进行复核。",
    ),
)


CHAPTER_THREE_BEFORE_TABLE: tuple[tuple[str, str], ...] = (
    (
        "",
        "预期成果包括一套通信受限条件下唯一任务关系保持与安全恢复方法、一套可复现的网络故障注入条件，以及覆盖直接推举、分区保持、再次推举和恢复交还的试验记录。结果以日志可计算指标、重复试验统计和失败明细为依据，安全性与时效性分别判定。表1列出阶段验收指标，时延目标先按状态交换周期计算并同时记录秒数；首轮基线可根据实测情况确定时延参数，双重发布、无效计划误接收和少数侧新增任务三类安全指标保持零容忍。",
    ),
    (
        "",
        "唯一任务关系保持率以全部生效任务关系为分母，逐条检查发布者、接管轮次、计划版本、有效期、目标和执行成员是否一致。接管时延从中心失效确认成立起算，到首份取得必要成员确认并进入有效期的临时计划为止，临时中心机当选但没有有效计划时仍记为未完成接管。恢复时延从原中心首次提供新鲜摘要起算，到更高轮次交还计划生效为止，并单列并行核对与冲突暂停时间；拒绝、超时、保持和失败样本均进入统计。",
    ),
)


INDICATOR_ROWS: tuple[tuple[str, str, str], ...] = (
    (
        "唯一任务关系保持率",
        "100%",
        "逐条核对已生效任务，检查发布者、轮次、版本、有效期、目标和执行成员；出现一条重复或矛盾关系即不通过。",
    ),
    (
        "双重临时中心生效次数",
        "0次",
        "按接管轮次核对全部过半确认与发布记录；同一轮出现两个有效发布者即不通过。",
    ),
    (
        "旧、冲突或无法核验计划误接收数",
        "0条",
        "主动重放旧消息、构造同版本不同内容并删除必要字段，核对各机接收与执行日志。",
    ),
    (
        "分区少数侧新增任务数",
        "0项",
        "按固定成员清单重建分区人数，检查未过半一侧是否只维持分区前已确认且安全的任务。",
    ),
    (
        "中心失效确认时延",
        "百分之九十五分位不超过3个交换周期",
        "从最后一条新鲜中心消息超过允许更新时间起算，到过半成员形成一致失效判断为止。",
    ),
    (
        "临时计划接管时延",
        "百分之九十五分位不超过8个交换周期",
        "从失效判断成立起算，到首份通过过半推举、版本和有效期检查的计划生效为止。",
    ),
    (
        "冲突未消解时错误交还次数",
        "0次",
        "中心恢复后至少连续5个核对窗口一致才允许交还；任一窗口存在冲突时检查是否继续降级保持。",
    ),
)


CHAPTER_THREE_AFTER_TABLE: tuple[tuple[str, str], ...] = (
    (
        "",
        "第一阶段逐项注入延迟、丢包、断链和网络分区，并叠加组合故障。延迟档设置为0、50、200、500和1000毫秒并附加随机抖动，丢包档设置为0%、5%、10%、30%和50%，断链时长设置为1、3、5、10和20秒，分区条件覆盖明确多数、人数相等及成员反复进出。每种单故障计划不少于100次随机时刻注入，组合故障不少于50次，完整记录故障起点、持续时间、受影响链路和消息类型，极端样本单独报告。",
    ),
    (
        "",
        "故障注入达到准入条件后，第二阶段在真实通信设备与飞控计算单元组成的台架闭环中检查队列积压、消息乱序和时钟偏差。试验先以固定位置复现单一故障，再通过移动或遮挡节点形成可控分区，机上日志和独立记录设备按同一时源离线对齐。唯一关系、双重发布、无效计划拒绝和少数侧限制四项安全指标全部通过后进入室内飞行；时延未达标时允许调整交换周期、合并窗口和排序参数，但等待窗口、过半门槛及有效期检查保持不变，参数变更后重跑基线与关键故障。",
    ),
    (
        "",
        "第三阶段依次开展2对2、4对4、6对6室内测试，每个规模先完成无故障基线，再注入延迟、丢包、中心断链、临时中心再次失联和网络分区。2对2检查一比一分区时不自行推举，4对4检查三比一条件下的接管和少数侧保持，6对6检查四比二接管、三比三无法过半及消息负荷。每类关键故障计划不少于20次独立重复，故障分别发生在任务初段、中段和计划临近到期时；飞行速度、高度、隔离区及安全动作经专项风险评估确定，无法核验的任务执行预定保持或退出。",
    ),
    (
        "",
        "每次试验记录消息发送与到达时刻、来源、序号、消息种类、固定成员清单校验值、接管轮次、计划版本、生成时刻、有效期、确认成员、拒绝原因、任务和执行状态。离线按相同规则重建各时刻的有效发布者与任务关系，并与机上判断逐条对照，日志不完整的样本从通过性统计中剔除并单列原因。结果分别给出样本数、中位数、百分之九十五分位数、最大值和失败明细，无故障、单故障、组合故障及不同规模分开报告，全部安全指标满足表1且时延达到预先确定的目标后形成阶段结论。",
    ),
    (
        "",
        "当前阶段验证结论仅适用于室内可控通信、规定规模及设定故障。远距离覆盖、复杂电磁对抗、长时运行、消息认证、时钟漂移和大范围分区须另行试验。双重发布、错误接收或少数侧新增任务未查明原因前，不得扩大试验范围。",
    ),
)


def configure_matplotlib() -> None:
    if FIGURE_FONT_PATH.exists():
        font_manager.fontManager.addfont(str(FIGURE_FONT_PATH))
        figure_font = font_manager.FontProperties(fname=str(FIGURE_FONT_PATH)).get_name()
    else:
        figure_font = "Noto Sans CJK SC"
    matplotlib.rcParams.update(
        {
            "font.family": figure_font,
            "font.sans-serif": [figure_font, "Noto Sans CJK SC", "DejaVu Sans"],
            "axes.unicode_minus": False,
            "figure.facecolor": "white",
            "text.color": INK,
        }
    )
    ASSET_DIR.mkdir(parents=True, exist_ok=True)


def new_canvas(title: str, subtitle: str = "", *, height: float = 8.0):
    figure, axis = plt.subplots(figsize=(14, height))
    figure.subplots_adjust(left=0, right=1, top=1, bottom=0)
    axis.set_xlim(0, 14)
    axis.set_ylim(0, height)
    axis.axis("off")
    axis.text(7, height - 0.42, title, ha="center", va="center", fontsize=29, weight="bold", color=INK)
    if subtitle:
        axis.text(7, height - 0.88, subtitle, ha="center", va="center", fontsize=18, color=MUTED)
    return figure, axis


def rounded_box(
    axis,
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    *,
    edge: str = LINE,
    face: str = LIGHT,
    fontsize: float = 18,
    weight: str = "normal",
    text_color: str = INK,
    radius: float = 0.08,
    linewidth: float = 1.8,
) -> None:
    axis.add_patch(
        FancyBboxPatch(
            (x, y),
            width,
            height,
            boxstyle=f"round,pad=0.025,rounding_size={radius}",
            facecolor=face,
            edgecolor=edge,
            linewidth=linewidth,
        )
    )
    axis.text(
        x + width / 2,
        y + height / 2,
        text,
        ha="center",
        va="center",
        fontsize=fontsize,
        weight=weight,
        color=text_color,
        linespacing=1.35,
    )


def arrow(
    axis,
    start: tuple[float, float],
    end: tuple[float, float],
    *,
    color: str = MUTED,
    linewidth: float = 2.0,
    dashed: bool = False,
    connectionstyle: str = "arc3",
) -> None:
    axis.annotate(
        "",
        xy=end,
        xytext=start,
        arrowprops={
            "arrowstyle": "-|>",
            "color": color,
            "lw": linewidth,
            "linestyle": "--" if dashed else "-",
            "mutation_scale": 17,
            "connectionstyle": connectionstyle,
        },
    )


def line(
    axis,
    start: tuple[float, float],
    end: tuple[float, float],
    *,
    color: str = MUTED,
    linewidth: float = 1.6,
    dashed: bool = False,
) -> None:
    axis.plot(
        [start[0], end[0]],
        [start[1], end[1]],
        color=color,
        lw=linewidth,
        linestyle="--" if dashed else "-",
        solid_capstyle="round",
    )


def drone(axis, x: float, y: float, label: str, *, color: str = BLUE, selected: bool = False) -> None:
    ring = 0.34 if selected else 0.29
    if selected:
        axis.add_patch(Circle((x, y), ring + 0.11, facecolor=PALE_ORANGE, edgecolor=ORANGE, lw=2.0))
    axis.add_patch(Circle((x, y), ring, facecolor="white", edgecolor=color, lw=2.2))
    axis.add_patch(Circle((x, y), 0.10, facecolor=color, edgecolor=color, lw=1.0))
    for dx, dy in ((-0.34, -0.23), (-0.34, 0.23), (0.34, -0.23), (0.34, 0.23)):
        line(axis, (x, y), (x + dx * 0.72, y + dy * 0.72), color=color, linewidth=2.0)
        axis.add_patch(Circle((x + dx, y + dy), 0.075, facecolor="white", edgecolor=color, lw=1.7))
    axis.text(x, y - 0.55, label, ha="center", va="top", fontsize=16.5, color=color, weight="bold")


def target(axis, x: float, y: float, label: str) -> None:
    axis.add_patch(
        Polygon(
            ((x, y + 0.23), (x - 0.24, y - 0.18), (x + 0.24, y - 0.18)),
            closed=True,
            facecolor=PALE_RED,
            edgecolor=RED,
            lw=2.0,
        )
    )
    axis.text(x, y - 0.35, label, ha="center", va="top", fontsize=15.5, color=RED, weight="bold")


def save_figure(figure, filename: str) -> None:
    path = ASSET_DIR / filename
    figure.savefig(path, dpi=300, facecolor="white")
    plt.close(figure)


def draw_communication_interruption_scenario() -> None:
    figure, axis = new_canvas(
        "通信中断条件下的任务场景",
        "无二级节点：中心失效后，由拦截无人机直接推举临时中心机",
    )
    rounded_box(
        axis,
        5.35,
        6.0,
        3.3,
        0.82,
        "中心节点\n消息停止更新",
        edge=RED,
        face=PALE_RED,
        fontsize=19,
        weight="bold",
        text_color=RED,
    )
    axis.text(8.82, 6.4, "×", ha="center", va="center", fontsize=39, color=RED, weight="bold")
    rounded_box(
        axis,
        0.75,
        5.15,
        12.5,
        0.62,
        "可能同时出现：消息延迟    随机丢包    持续断链    网络分区",
        edge=LINE,
        face=LIGHT,
        fontsize=18,
        text_color=INK,
    )
    for x in (2.0, 4.05, 6.1, 7.85, 10.15, 12.25):
        arrow(axis, (7.0, 5.99), (x, 4.58), color=RED, linewidth=1.5, dashed=True)

    axis.add_patch(Rectangle((8.55, 1.1), 0.28, 3.75, facecolor=PALE_RED, edgecolor=RED, lw=1.5, alpha=0.85))
    axis.text(8.69, 4.55, "通信分区", ha="center", va="center", fontsize=16, color=RED, weight="bold", rotation=90)

    left_nodes = ((1.35, 3.75), (3.25, 3.05), (5.2, 3.75), (7.15, 3.05))
    right_nodes = ((10.05, 3.65), (12.35, 3.05))
    for idx, (x, y) in enumerate(left_nodes, start=1):
        drone(axis, x, y, f"无人机{idx}", color=BLUE, selected=idx == 2)
    for idx, (x, y) in enumerate(right_nodes, start=5):
        drone(axis, x, y, f"无人机{idx}", color=TEAL)

    for i, start in enumerate(left_nodes):
        for end in left_nodes[i + 1 :]:
            line(axis, start, end, color="#9AB2C2", linewidth=1.3, dashed=True)
    line(axis, right_nodes[0], right_nodes[1], color="#8FB8B3", linewidth=1.5, dashed=True)

    target(axis, 2.25, 1.35, "目标1")
    target(axis, 5.65, 1.35, "目标2")
    target(axis, 10.25, 1.35, "目标3")
    target(axis, 12.7, 1.35, "目标4")
    arrow(axis, (1.55, 3.38), (2.15, 1.68), color=GREEN, linewidth=2.5)
    arrow(axis, (5.05, 3.38), (5.6, 1.68), color=GREEN, linewidth=2.5)
    arrow(axis, (10.08, 3.27), (10.23, 1.68), color=GREEN, linewidth=2.5)
    arrow(axis, (12.38, 2.67), (12.68, 1.68), color=GREEN, linewidth=2.5)

    rounded_box(
        axis,
        0.55,
        0.14,
        7.5,
        0.62,
        "多数连通侧：取得原成员清单过半确认后，方可推举临时中心机并调整可核验任务",
        edge=BLUE,
        face=PALE_BLUE,
        fontsize=16.5,
        text_color=BLUE,
    )
    rounded_box(
        axis,
        8.95,
        0.14,
        4.5,
        0.62,
        "少数侧：只维持已确认且安全的任务\n不得产生新分配",
        edge=TEAL,
        face=PALE_TEAL,
        fontsize=16.5,
        text_color=TEAL,
    )
    save_figure(figure, FIGURES[0][0])


def draw_temporary_center_election_principle() -> None:
    figure, axis = new_canvas(
        "临时中心机推举原理",
        "固定成员清单、固定排序、编号消解、同轮单次确认、超过半数生效",
    )
    rounded_box(
        axis,
        0.55,
        6.12,
        3.0,
        0.9,
        "共同输入\n成员清单\n新鲜状态摘要",
        edge=BLUE,
        face=PALE_BLUE,
        fontsize=16.5,
        weight="bold",
    )
    rounded_box(
        axis,
        4.12,
        6.12,
        3.0,
        0.9,
        "硬条件筛选\n通信、航时、计算\n任务安全",
        edge=TEAL,
        face=PALE_TEAL,
        fontsize=16.5,
        weight="bold",
    )
    rounded_box(
        axis,
        7.45,
        6.12,
        3.0,
        0.9,
        "固定候选排序\n覆盖、航时、计算\n位置、负荷",
        edge=ORANGE,
        face=PALE_ORANGE,
        fontsize=16.5,
        weight="bold",
    )
    rounded_box(
        axis,
        10.78,
        6.12,
        2.68,
        0.9,
        "同档消解\n按无人机编号\n确定先后",
        edge=RED,
        face=PALE_RED,
        fontsize=16.5,
        weight="bold",
    )
    arrow(axis, (3.8, 6.52), (4.1, 6.52), color=BLUE)
    arrow(axis, (7.12, 6.52), (7.43, 6.52), color=TEAL)
    arrow(axis, (10.45, 6.52), (10.76, 6.52), color=ORANGE)

    rounded_box(
        axis,
        0.7,
        4.55,
        2.8,
        0.92,
        "候选一\n覆盖优  航时中  负荷高",
        edge=LINE,
        face="white",
        fontsize=17,
    )
    rounded_box(
        axis,
        3.65,
        4.55,
        2.8,
        0.92,
        "候选二\n覆盖优  航时足  负荷低",
        edge=ORANGE,
        face=PALE_ORANGE,
        fontsize=17,
        weight="bold",
        text_color=ORANGE,
        linewidth=2.5,
    )
    rounded_box(
        axis,
        6.6,
        4.55,
        2.8,
        0.92,
        "候选三\n覆盖中  航时足  位置偏",
        edge=LINE,
        face="white",
        fontsize=17,
    )
    rounded_box(
        axis,
        9.55,
        4.55,
        2.8,
        0.92,
        "候选四\n覆盖不足  不进入推举",
        edge=RED,
        face=PALE_RED,
        fontsize=17,
        text_color=RED,
    )
    axis.text(5.05, 5.73, "固定结果：候选二排序首位", ha="center", va="center", fontsize=16.5, color=ORANGE, weight="bold")

    vote_nodes = ((1.25, 3.05), (3.0, 3.05), (4.75, 3.05), (6.5, 3.05), (8.25, 3.05), (10.0, 3.05))
    for idx, (x, y) in enumerate(vote_nodes, start=1):
        drone(axis, x, y, str(idx), color=BLUE if idx <= 4 else MUTED)
        if idx <= 4:
            arrow(axis, (x, y - 0.15), (6.55, 2.07), color=BLUE, linewidth=1.8)
        else:
            arrow(axis, (x, y - 0.15), (11.4, 1.85), color=MUTED, linewidth=1.5, dashed=True)

    rounded_box(
        axis,
        4.95,
        0.9,
        3.2,
        1.08,
        "候选二取得4份确认\n6名成员中超过半数\n限时生效",
        edge=ORANGE,
        face=PALE_ORANGE,
        fontsize=16.5,
        weight="bold",
        text_color=ORANGE,
        linewidth=2.5,
    )
    rounded_box(
        axis,
        10.0,
        0.95,
        2.8,
        0.98,
        "其他候选确认不足\n不得发布新任务",
        edge=LINE,
        face=LIGHT,
        fontsize=17,
        text_color=MUTED,
    )
    rounded_box(
        axis,
        0.55,
        0.12,
        12.9,
        0.56,
        "同一固定成员清单 + 同一接管轮次 + 每名成员只确认一名候选 + 过半门槛  →  合法发布者最多一个",
        edge=GREEN,
        face=PALE_GREEN,
        fontsize=18,
        weight="bold",
        text_color=GREEN,
    )
    save_figure(figure, FIGURES[1][0])


def draw_degradation_recovery_full_flow() -> None:
    height = 9.0
    figure, axis = new_canvas(
        "降级与恢复完整流程",
        "状态转移均核对消息时效、唯一发布者、过半确认、版本和有效期",
        height=height,
    )
    axis.add_patch(Rectangle((0.25, 6.42), 13.5, 1.9, facecolor="#F8FAFB", edgecolor=LINE, lw=1.0))
    axis.add_patch(Rectangle((0.25, 3.22), 13.5, 2.8, facecolor="#FCFAF8", edgecolor=LINE, lw=1.0))
    axis.add_patch(Rectangle((0.25, 0.32), 13.5, 2.42, facecolor="#F8FBF9", edgecolor=LINE, lw=1.0))

    top_boxes = (
        (0.75, "中心正常\n低频刷新、事件触发", BLUE, PALE_BLUE),
        (3.45, "消息更新时间超限\n进入可疑观察", ORANGE, PALE_ORANGE),
        (6.15, "持续窗口达到门槛\n过半成员确认失效", RED, PALE_RED),
        (8.85, "固定共同成员清单\n旧消息不得倒退轮次", TEAL, PALE_TEAL),
        (11.55, "接管轮次加一\n启动候选排序", ORANGE, PALE_ORANGE),
    )
    for x, text, edge, face in top_boxes:
        rounded_box(axis, x, 6.73, 2.15, 1.15, text, edge=edge, face=face, fontsize=17, weight="bold")
    for left, right in zip(top_boxes, top_boxes[1:]):
        arrow(axis, (left[0] + 2.15, 7.3), (right[0], 7.3), color=MUTED, linewidth=1.8)

    middle_boxes = (
        (0.75, 4.62, "固定排序\n同档按编号", BLUE, PALE_BLUE),
        (3.25, 4.62, "过半确认\n每机每轮一次", ORANGE, PALE_ORANGE),
        (5.75, 4.62, "临时中心机\n取得限时发布权", GREEN, PALE_GREEN),
        (8.25, 4.62, "计划带轮次、版本\n生成时刻和有效期", TEAL, PALE_TEAL),
        (10.75, 4.62, "必要成员确认\n任务关系方可生效", BLUE, PALE_BLUE),
    )
    for x, y, text, edge, face in middle_boxes:
        rounded_box(axis, x, y, 2.15, 0.9, text, edge=edge, face=face, fontsize=16.8, weight="bold")
    arrow(axis, (12.62, 6.72), (1.82, 5.54), color=ORANGE, linewidth=2.0, connectionstyle="arc3,rad=0.12")
    for left, right in zip(middle_boxes, middle_boxes[1:]):
        arrow(axis, (left[0] + 2.15, 5.07), (right[0], 5.07), color=MUTED, linewidth=1.7)

    rounded_box(axis, 0.75, 3.38, 3.25, 0.86, "推举超时\n保持安全任务，进入更高轮次重选", edge=RED, face=PALE_RED, fontsize=14.5, text_color=RED, weight="bold")
    rounded_box(axis, 4.25, 3.38, 3.25, 0.86, "临时中心再次失联\n授权到期，停止新发布", edge=RED, face=PALE_RED, fontsize=14.5, text_color=RED, weight="bold")
    rounded_box(axis, 7.75, 3.38, 5.0, 0.86, "网络分区\n多数侧只调整可核验任务；少数侧不重新分配", edge=ORANGE, face=PALE_ORANGE, fontsize=14.5, text_color=ORANGE, weight="bold")
    arrow(axis, (4.32, 4.62), (2.38, 4.2), color=RED, dashed=True, connectionstyle="arc3,rad=0.18")
    arrow(axis, (6.82, 4.62), (5.88, 4.2), color=RED, dashed=True)
    arrow(axis, (9.32, 4.62), (10.25, 4.2), color=ORANGE, dashed=True)

    recovery_boxes = (
        (0.75, "原中心恢复消息\n先作为观察者", GREEN, PALE_GREEN),
        (3.45, "并行核对\n航迹、资源、计划\n执行状态", TEAL, PALE_TEAL),
        (6.15, "状态是否一致？", ORANGE, PALE_ORANGE),
        (8.85, "过半确认更高轮次\n明确时刻交还", GREEN, PALE_GREEN),
        (11.55, "临时中心停止续期\n中心新版本生效", BLUE, PALE_BLUE),
    )
    for x, text, edge, face in recovery_boxes:
        rounded_box(axis, x, 1.2, 2.15, 1.0, text, edge=edge, face=face, fontsize=16.8, weight="bold")
    for left, right in zip(recovery_boxes, recovery_boxes[1:]):
        arrow(axis, (left[0] + 2.15, 1.7), (right[0], 1.7), color=MUTED, linewidth=1.7)
    axis.text(8.55, 1.97, "一致", ha="center", va="center", fontsize=14.5, color=GREEN, weight="bold")
    rounded_box(axis, 5.55, 0.42, 3.35, 0.48, "不一致：继续降级，冲突任务保持", edge=RED, face=PALE_RED, fontsize=13.8, text_color=RED, weight="bold")
    arrow(axis, (7.22, 1.2), (7.22, 0.9), color=RED, dashed=True)
    save_figure(figure, FIGURES[2][0])


def draw_network_fault_injection_route() -> None:
    figure, axis = new_canvas(
        "网络故障注入实施路线",
        "先验证安全规则和时序边界，再开展2对2、4对4、6对6室内测试",
    )
    rounded_box(
        axis,
        0.55,
        5.75,
        2.25,
        1.08,
        "第一阶段\n可重复故障注入",
        edge=BLUE,
        face=PALE_BLUE,
        fontsize=19,
        weight="bold",
        text_color=BLUE,
    )
    fault_boxes = (
        (3.15, "延迟阶梯\n含随机抖动", BLUE, PALE_BLUE),
        (5.65, "丢包比例\n含关键确认丢失", ORANGE, PALE_ORANGE),
        (8.15, "断链时长\n含恢复后旧包到达", RED, PALE_RED),
        (10.65, "网络分区\n多数、等分、反复变化", TEAL, PALE_TEAL),
    )
    for x, text, edge, face in fault_boxes:
        rounded_box(axis, x, 5.75, 2.25, 1.08, text, edge=edge, face=face, fontsize=17.5, weight="bold")
    arrow(axis, (2.8, 6.29), (3.13, 6.29), color=BLUE)
    for left, right in zip(fault_boxes, fault_boxes[1:]):
        arrow(axis, (left[0] + 2.25, 6.29), (right[0], 6.29), color=MUTED, linewidth=1.5)

    rounded_box(
        axis,
        1.0,
        4.15,
        12.0,
        0.86,
        "采集字段：发送与到达时刻、来源、序号、接管轮次、计划版本、有效期、确认成员、拒绝原因、任务状态",
        edge=LINE,
        face=LIGHT,
        fontsize=17,
    )
    for x in (1.68, 4.28, 6.78, 9.28, 11.78):
        arrow(axis, (x, 5.75), (x, 5.03), color=MUTED, linewidth=1.4, dashed=True)

    metric_boxes = (
        (0.75, "唯一关系\n保持率", GREEN, PALE_GREEN),
        (3.4, "双重来源\n生效次数", RED, PALE_RED),
        (6.05, "无效计划\n误接收数", ORANGE, PALE_ORANGE),
        (8.7, "接管与恢复\n时延分布", BLUE, PALE_BLUE),
        (11.35, "少数侧\n新增任务数", TEAL, PALE_TEAL),
    )
    for x, text, edge, face in metric_boxes:
        rounded_box(axis, x, 2.95, 1.9, 0.82, text, edge=edge, face=face, fontsize=17, weight="bold")
    arrow(axis, (7.0, 4.15), (7.0, 3.78), color=MUTED)

    stages = (
        (0.55, "台架闭环\n固定位置与可控遮挡", BLUE, PALE_BLUE),
        (3.25, "2对2室内\n等分时不自行推举", TEAL, PALE_TEAL),
        (5.95, "4对4室内\n三比一接管与保持", GREEN, PALE_GREEN),
        (8.65, "6对6室内\n四比二与三比三", ORANGE, PALE_ORANGE),
        (11.35, "阶段评审\n保留条件与限制", RED, PALE_RED),
    )
    for x, text, edge, face in stages:
        rounded_box(axis, x, 1.2, 2.1, 0.95, text, edge=edge, face=face, fontsize=16.7, weight="bold")
    for left, right in zip(stages, stages[1:]):
        arrow(axis, (left[0] + 2.1, 1.68), (right[0], 1.68), color=MUTED, linewidth=1.6)
    rounded_box(
        axis,
        1.2,
        0.18,
        11.6,
        0.58,
        "逐级准入：安全指标全部通过后方可进入下一阶段；时延优化不得降低过半确认和有效期检查",
        edge=GREEN,
        face=PALE_GREEN,
        fontsize=17.5,
        weight="bold",
        text_color=GREEN,
    )
    save_figure(figure, FIGURES[3][0])


def build_figures() -> None:
    configure_matplotlib()
    draw_communication_interruption_scenario()
    draw_temporary_center_election_principle()
    draw_degradation_recovery_full_flow()
    draw_network_fault_injection_route()


def set_run_font(
    run,
    *,
    size: float,
    bold: bool = False,
    color: str = WORD_INK,
    east_asia: str = BODY_FONT,
    latin: str = LATIN_FONT,
) -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_row_cant_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_repeat_table_header(row) -> None:
    repeat_table_header(row)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    field = OxmlElement("w:instrText")
    field.set(qn("xml:space"), "preserve")
    field.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, field, separate, placeholder, end))
    set_run_font(run, size=9.5, color=WORD_MUTED)


def configure_section(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.35)
    section.bottom_margin = Cm(2.25)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.05)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.98)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    body_style = styles.add_style("Report Body", 1)
    body_style.base_style = normal
    body_style.font.name = LATIN_FONT
    body_style.font.size = Pt(14)
    body_style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    body_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body_style.paragraph_format.first_line_indent = Cm(0.98)
    body_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    body_style.paragraph_format.space_after = Pt(5)
    body_style.paragraph_format.widow_control = True

    caption_style = styles.add_style("Report Caption", 1)
    caption_style.base_style = normal
    caption_style.font.name = LATIN_FONT
    caption_style.font.size = Pt(10.5)
    caption_style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_style.paragraph_format.first_line_indent = Cm(0)
    caption_style.paragraph_format.line_spacing = 1.15
    caption_style.paragraph_format.space_before = Pt(2)
    caption_style.paragraph_format.space_after = Pt(8)
    caption_style.paragraph_format.keep_with_next = False

    heading = styles["Heading 1"]
    heading.font.name = LATIN_FONT
    heading.font.size = Pt(18)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(WORD_BLUE)
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
    heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading.paragraph_format.first_line_indent = Cm(0)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(11)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.widow_control = True


def configure_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(
        header.add_run("通信受限条件下一致分配与降级接管"),
        size=9.5,
        color=WORD_MUTED,
        east_asia=BODY_FONT,
    )
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), WORD_LINE)
    border.append(bottom)
    header._p.get_or_add_pPr().append(border)

    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    set_run_font(footer.add_run("专题汇报材料  |  第 "), size=9.5, color=WORD_MUTED)
    add_field(footer, "PAGE")
    set_run_font(footer.add_run(" 页"), size=9.5, color=WORD_MUTED)

    section_properties = section._sectPr
    page_number_type = section_properties.find(qn("w:pgNumType"))
    if page_number_type is None:
        page_number_type = OxmlElement("w:pgNumType")
        section_properties.append(page_number_type)
    page_number_type.set(qn("w:start"), "1")


def add_cover(document: Document) -> None:
    for _ in range(4):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.space_after = Pt(18)
    set_run_font(title.add_run(TITLE), size=26, bold=True, color=WORD_INK, east_asia=HEADING_FONT)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Cm(0)
    subtitle.paragraph_format.space_after = Pt(8)
    set_run_font(
        subtitle.add_run("中心失效后由拦截无人机直接推举临时中心机"),
        size=16,
        bold=True,
        color=WORD_BLUE,
        east_asia=HEADING_FONT,
    )
    boundary = document.add_paragraph()
    boundary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    boundary.paragraph_format.first_line_indent = Cm(0)
    set_run_font(boundary.add_run("科学问题与拟实施方案"), size=13, color=WORD_TEAL, east_asia=HEADING_FONT)

    for _ in range(10):
        document.add_paragraph()
    owner = document.add_paragraph()
    owner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    owner.paragraph_format.first_line_indent = Cm(0)
    set_run_font(owner.add_run("项目组"), size=13, bold=True, color=WORD_INK)
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.first_line_indent = Cm(0)
    set_run_font(date.add_run("2026 年 8 月"), size=12, color=WORD_MUTED)


def add_chapter_heading(document: Document, title: str, *, page_break: bool) -> None:
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.page_break_before = page_break
    set_run_font(paragraph.add_run(title), size=18, bold=True, color=WORD_BLUE, east_asia=HEADING_FONT)


def add_body_paragraph(document: Document, lead: str, body: str) -> None:
    paragraph = document.add_paragraph(style="Report Body")
    if lead:
        set_run_font(paragraph.add_run(lead), size=14, bold=True, color=WORD_INK, east_asia=HEADING_FONT)
    set_run_font(paragraph.add_run(body), size=14, color=WORD_INK, east_asia=BODY_FONT)


def add_body_paragraphs(document: Document, paragraphs: Iterable[tuple[str, str]]) -> None:
    for lead, body in paragraphs:
        add_body_paragraph(document, lead, body)


def add_figure(document: Document, filename: str, caption: str, number: int) -> None:
    image_path = ASSET_DIR / filename
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.first_line_indent = Cm(0)
    picture_paragraph.paragraph_format.space_before = Pt(5)
    picture_paragraph.paragraph_format.space_after = Pt(0)
    picture_paragraph.paragraph_format.keep_together = True
    picture_paragraph.paragraph_format.keep_with_next = True
    picture_paragraph.add_run().add_picture(str(image_path), width=Cm(16.0))

    caption_paragraph = document.add_paragraph(style="Report Caption")
    set_run_font(
        caption_paragraph.add_run(f"图 {number}  {caption}"),
        size=10.5,
        color=WORD_MUTED,
        east_asia=BODY_FONT,
    )


def add_indicator_table(document: Document) -> None:
    caption = document.add_paragraph(style="Report Caption")
    caption.paragraph_format.keep_with_next = True
    set_run_font(caption.add_run("表 1  阶段验收指标与判定方法"), size=10.5, color=WORD_MUTED)

    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    table.columns[0].width = Cm(3.45)
    table.columns[1].width = Cm(3.65)
    table.columns[2].width = Cm(8.9)
    headers = ("指标", "阶段目标", "检验与判定方法")
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    set_row_cant_split(header_row)
    for cell, text in zip(header_row.cells, headers):
        set_cell_shading(cell, WORD_BLUE)
        set_cell_margins(cell, top=110, bottom=110, start=110, end=110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.1
        set_run_font(paragraph.add_run(text), size=11, bold=True, color="FFFFFF", east_asia=HEADING_FONT)

    for row_index, row_data in enumerate(INDICATOR_ROWS):
        row = table.add_row()
        set_row_cant_split(row)
        for column_index, (cell, text) in enumerate(zip(row.cells, row_data)):
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else "F4F7F9")
            set_cell_margins(cell, top=90, bottom=90, start=100, end=100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index < 2 else WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            set_run_font(
                paragraph.add_run(text),
                size=10.5,
                bold=column_index == 1,
                color=WORD_INK if column_index != 1 else WORD_TEAL,
                east_asia=BODY_FONT,
            )


def build_document() -> None:
    document = Document()
    configure_section(document.sections[0])
    configure_styles(document)
    add_cover(document)

    content_section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(content_section)
    configure_header_footer(content_section)

    add_chapter_heading(document, CHAPTER_TITLES[0], page_break=False)
    add_body_paragraphs(document, CHAPTER_ONE[:2])
    add_body_paragraph(document, *CHAPTER_ONE[2])
    add_figure(document, *FIGURES[0], number=1)
    add_body_paragraphs(document, CHAPTER_ONE[3:])

    add_chapter_heading(document, CHAPTER_TITLES[1], page_break=True)
    add_body_paragraphs(document, CHAPTER_TWO_BEFORE_FIGURE_TWO)
    add_figure(document, *FIGURES[1], number=2)
    add_body_paragraphs(document, CHAPTER_TWO_AFTER_FIGURE_TWO[:-1])
    add_figure(document, *FIGURES[2], number=3)
    add_body_paragraph(document, *CHAPTER_TWO_AFTER_FIGURE_TWO[-1])

    add_chapter_heading(document, CHAPTER_TITLES[2], page_break=True)
    add_body_paragraphs(document, CHAPTER_THREE_BEFORE_TABLE)
    add_indicator_table(document)
    add_body_paragraph(document, *CHAPTER_THREE_AFTER_TABLE[0])
    add_figure(document, *FIGURES[3], number=4)
    add_body_paragraphs(document, CHAPTER_THREE_AFTER_TABLE[1:])

    properties = document.core_properties
    properties.title = TITLE
    properties.subject = "中心节点失效后拦截无人机直接推举临时中心机的科学问题与拟实施方案"
    properties.author = "项目组"
    properties.keywords = "通信受限, 一致分配, 临时中心机, 降级接管, 安全恢复"
    properties.comments = "直接推举与室内验证尚待实施，现有基础限于版本、时效、拒绝和故障注入。"
    document.save(OUTPUT)


def iter_document_text(document: Document) -> Iterable[str]:
    for paragraph in document.paragraphs:
        if paragraph.text:
            yield paragraph.text
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text:
                        yield paragraph.text


def validate_figures() -> dict[str, tuple[int, int]]:
    dimensions: dict[str, tuple[int, int]] = {}
    for filename, _ in FIGURES:
        path = ASSET_DIR / filename
        if not path.exists():
            raise RuntimeError(f"missing figure: {path}")
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            width, height = image.size
        if width < 4000 or height < 2200:
            raise RuntimeError(f"figure resolution too low: {filename}={width}x{height}")
        dimensions[filename] = (width, height)
    return dimensions


def validate_document() -> dict[str, int]:
    if not OUTPUT.exists():
        raise RuntimeError(f"missing output: {OUTPUT}")
    document = Document(OUTPUT)
    with ZipFile(OUTPUT) as archive:
        damaged = archive.testzip()
        if damaged is not None:
            raise RuntimeError(f"damaged DOCX member: {damaged}")
        media = [name for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/")]
        required_parts = {"[Content_Types].xml", "word/document.xml", "word/styles.xml", "word/_rels/document.xml.rels"}
        missing_parts = required_parts - set(archive.namelist())
        if missing_parts:
            raise RuntimeError(f"missing DOCX parts: {sorted(missing_parts)}")
    if len(media) != len(FIGURES):
        raise RuntimeError(f"expected {len(FIGURES)} embedded images, found {len(media)}")
    if len(document.inline_shapes) != len(FIGURES):
        raise RuntimeError(f"expected {len(FIGURES)} inline shapes, found {len(document.inline_shapes)}")

    headings = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"]
    if tuple(headings) != CHAPTER_TITLES:
        raise RuntimeError(f"chapter structure changed: {headings}")
    extra_headings = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading ") and paragraph.style.name != "Heading 1"
    ]
    if extra_headings:
        raise RuntimeError(f"unexpected subordinate headings: {extra_headings}")

    all_text = "\n".join(iter_document_text(document))
    han_count = len(re.findall(r"[\u3400-\u4dbf\u4e00-\u9fff]", all_text))
    if not 4000 <= han_count <= 6000:
        raise RuntimeError(f"Chinese character count out of range: {han_count}")
    for required in (
        TITLE,
        "针对上述问题，拟重点突破“时变通信条件下唯一任务关系保持与安全恢复技术”。该技术",
        "体系不设置二级节点",
        "事件触发交换",
        "通信覆盖、剩余航时、计算能力、位置条件和任务负荷",
        "发布者编号、接管轮次、计划版本、生成时刻",
        "少数侧不能缩小成员清单",
        "再次排序",
        "临时中心机停止更新",
        "恢复交还",
        "方案中的直接推举和室内验证尚待实施；现有版本、时效、拒绝和故障注入基础可复用",
        "2对2、4对4、6对6室内测试",
    ):
        if required not in all_text:
            raise RuntimeError(f"missing required content: {required}")
    sensitive_phrases = (
        "本报告聚焦",
        "本报告面向",
        "本报告采用",
        "本报告将",
        "本报告不",
        "问题的实质不是",
        "核心不是",
        "通俗地说",
        "换言之",
        "写成",
        "统一表述为",
        "贯通",
        "收口",
        "赋能",
        "打造",
    )
    for forbidden in sensitive_phrases + (
        "合同",
        "槽位",
        "CBBA",
        "D4",
        "人工智能",
        "AI口号",
        "代码路径",
        "开发记录",
    ):
        if forbidden in all_text:
            raise RuntimeError(f"forbidden report wording: {forbidden}")
    completed_claims = re.findall(r"(?:已经|已)(?:完成|实现|达到|通过|验证|取得)", all_text)
    if completed_claims:
        raise RuntimeError(f"report contains completed-result wording: {completed_claims}")

    body_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Report Body"]
    if len(body_paragraphs) < 20:
        raise RuntimeError(f"too few developed body paragraphs: {len(body_paragraphs)}")
    short_body = []
    for paragraph in body_paragraphs:
        sentence_count = len(re.findall(r"[。！？]", paragraph.text))
        han = len(re.findall(r"[\u3400-\u4dbf\u4e00-\u9fff]", paragraph.text))
        if sentence_count < 3 or han < 90:
            short_body.append((paragraph.text[:40], sentence_count, han))
    if short_body:
        raise RuntimeError(f"body paragraphs are underdeveloped: {short_body}")

    if len(document.tables) != 1 or len(document.tables[0].rows) != len(INDICATOR_ROWS) + 1:
        raise RuntimeError("indicator table structure changed")
    return {
        "han_characters": han_count,
        "paragraphs": len(document.paragraphs),
        "body_paragraphs": len(body_paragraphs),
        "tables": len(document.tables),
        "images": len(media),
        "sensitive_phrases": sum(all_text.count(phrase) for phrase in sensitive_phrases),
        "bytes": OUTPUT.stat().st_size,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate-only", action="store_true", help="validate existing Word and PNG outputs")
    args = parser.parse_args()
    if not args.validate_only:
        build_figures()
        build_document()
    dimensions = validate_figures()
    metrics = validate_document()
    sizes = ", ".join(f"{name}={width}x{height}" for name, (width, height) in dimensions.items())
    print(
        f"{OUTPUT.name}: han={metrics['han_characters']}, paragraphs={metrics['paragraphs']}, "
        f"body_paragraphs={metrics['body_paragraphs']}, tables={metrics['tables']}, "
        f"images={metrics['images']}, sensitive={metrics['sensitive_phrases']}, bytes={metrics['bytes']}"
    )
    print(f"figures: {sizes}")


if __name__ == "__main__":
    main()
