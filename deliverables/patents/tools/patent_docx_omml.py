#!/usr/bin/env python3
"""Build patent Word drafts with native Word OMML equations.

Pandoc owns Markdown and TeX conversion.  python-docx is deliberately limited
to page and paragraph formatting so that the generated OMML nodes are kept.
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[3]
PATENT_DIR = ROOT / "deliverables" / "patents"
PATENTS = (
    PATENT_DIR / "基于多智能体决策的集群无人机协同反无人机系统_CN.md",
    PATENT_DIR / "基于图神经网络的双光电多目标轨迹配准与交会定位方法_CN.md",
)

MAIN_PAGE_BREAKS = {"摘要", "权利要求书", "说明书"}
CENTERED_HEADINGS = {"摘要", "摘要附图", "权利要求书", "说明书"}
DESCRIPTION_HEADINGS = {
    "技术领域",
    "背景技术",
    "发明内容",
    "附图说明",
    "具体实施方式",
    "附图标记",
}


def find_pandoc() -> Path:
    """Find Pandoc without requiring a system-wide installation."""
    candidates: list[str] = []
    if os.environ.get("PANDOC"):
        candidates.append(os.environ["PANDOC"])
    on_path = shutil.which("pandoc")
    if on_path:
        candidates.append(on_path)
    candidates.append(str(Path.home() / ".local" / "bin" / "pandoc"))
    for candidate in candidates:
        path = Path(candidate).expanduser()
        if path.is_file() and os.access(path, os.X_OK):
            return path.resolve()
    raise RuntimeError(
        "未找到Pandoc。请安装Pandoc，或通过PANDOC环境变量指定可执行文件。"
    )


def set_style_font(style, east_asia: str, latin: str, size: float, *, bold: bool = False) -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


def set_run_font(run, east_asia: str = "宋体", latin: str = "Times New Roman", size: float = 12) -> None:
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


def configure_reference(reference_path: Path) -> None:
    """Apply patent page and style defaults to Pandoc's reference document."""
    document = Document(reference_path)
    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.3)
        section.footer_distance = Cm(1.3)

    styles = document.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name not in styles:
            continue
        style = styles[name]
        set_style_font(style, "宋体", "Times New Roman", 12)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(24)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

    for name, east_asia, latin, size in (
        ("Title", "黑体", "Arial", 18),
        ("Heading 1", "黑体", "Arial", 16),
        ("Heading 2", "黑体", "Arial", 14),
        ("Heading 3", "黑体", "Arial", 12),
        ("Heading 4", "黑体", "Arial", 12),
    ):
        if name not in styles:
            continue
        style = styles[name]
        set_style_font(style, east_asia, latin, size, bold=True)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Caption" in styles:
        caption = styles["Caption"]
        set_style_font(caption, "宋体", "Times New Roman", 10.5)
        caption.font.italic = False
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(6)

    document.save(reference_path)


def create_reference_doc(pandoc: Path, reference_path: Path) -> None:
    with reference_path.open("wb") as stream:
        subprocess.run(
            [str(pandoc), "--print-default-data-file", "reference.docx"],
            stdout=stream,
            check=True,
        )
    configure_reference(reference_path)


def add_page_number(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        paragraph._p.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = paragraph.add_run("第 ")
    set_run_font(prefix, size=9)
    field_run = paragraph.add_run()
    set_run_font(field_run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend((begin, instruction, separate, placeholder, end))
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=9)


def configure_word_compatibility(document: Document) -> None:
    """Prevent justified hard-break lines from being stretched across the page."""
    settings = document.settings._element
    compatibility = settings.find(qn("w:compat"))
    if compatibility is None:
        compatibility = OxmlElement("w:compat")
        settings.append(compatibility)
    if compatibility.find(qn("w:doNotExpandShiftReturn")) is None:
        compatibility.append(OxmlElement("w:doNotExpandShiftReturn"))


def paragraph_after(paragraph: Paragraph) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    return Paragraph(element, paragraph._parent)


def document_style(document: Document, name: str):
    """Return a style by display name, including Pandoc built-in styles."""
    for style in document.styles:
        if style.name == name:
            return style
    raise KeyError(f"Word模板缺少样式：{name}")


def markdown_figure_captions(md_path: Path) -> list[str]:
    text = md_path.read_text(encoding="utf-8")
    return [alt.strip() for alt in re.findall(r"!\[([^]]*)]\([^)]+\)", text)]


def normalize_math_punctuation(text: str) -> str:
    """Use punctuation that Cambria Math and LibreOffice render consistently."""
    def normalize(match: re.Match[str]) -> str:
        opening, body, closing = match.groups()
        body = body.translate(str.maketrans({"，": ",", "。": ".", "；": ";", "：": ":"}))
        return f"{opening}{body}{closing}"

    display_pattern = re.compile(r"(?<!\\)(\$\$)(.*?)(?<!\\)(\$\$)", flags=re.DOTALL)
    text = display_pattern.sub(normalize, text)
    inline_pattern = re.compile(
        r"(?<![\\$])(\$)(?!\$)(.*?)(?<!\\)(\$)(?!\$)",
        flags=re.DOTALL,
    )
    return inline_pattern.sub(normalize, text)


def add_figure_captions(document: Document, captions: list[str]) -> None:
    figure_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph._p.xpath(".//w:drawing")
    ]
    if len(figure_paragraphs) != len(captions):
        raise RuntimeError(
            f"附图段落数与Markdown不一致：Word={len(figure_paragraphs)}，Markdown={len(captions)}"
        )
    for paragraph, caption_text in zip(figure_paragraphs, captions):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_with_next = True
        following = paragraph._p.getnext()
        if following is not None and following.tag == qn("w:p"):
            candidate = Paragraph(following, paragraph._parent)
            if candidate.text.strip() == caption_text:
                candidate.style = document_style(document, "Caption")
                continue
        caption = paragraph_after(paragraph)
        caption.style = document_style(document, "Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption.add_run(caption_text)
        set_run_font(run, size=10.5)


def size_images(document: Document) -> None:
    target_width = Cm(15.5)
    for shape in document.inline_shapes:
        if not shape.width:
            continue
        ratio = target_width / shape.width
        shape.height = int(shape.height * ratio)
        shape.width = target_width


def has_numbering(paragraph: Paragraph) -> bool:
    properties = paragraph._p.pPr
    return properties is not None and properties.numPr is not None


def restyle_paragraphs(document: Document, title: str) -> None:
    """Restyle paragraphs without altering text or OMML child nodes."""
    in_claims = False
    heading_map = {
        "Heading 2": "Heading 1",
        "Heading 3": "Heading 2",
        "Heading 4": "Heading 3",
    }

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        original_style = paragraph.style.name if paragraph.style else ""
        has_drawing = bool(paragraph._p.xpath(".//w:drawing"))
        has_display_math = bool(paragraph._p.xpath(".//m:oMathPara"))

        if text == title:
            paragraph.style = document_style(document, "Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(12)
            for run in paragraph.runs:
                set_run_font(run, "黑体", "Arial", 18)
                run.bold = True
            continue

        if text == "权利要求书":
            in_claims = True
        elif text == "说明书":
            in_claims = False

        if text in CENTERED_HEADINGS:
            paragraph.style = document_style(document, "Heading 1")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif text in DESCRIPTION_HEADINGS:
            paragraph.style = document_style(document, "Heading 1")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif original_style in heading_map:
            paragraph.style = document_style(document, heading_map[original_style])
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        if text in MAIN_PAGE_BREAKS:
            paragraph.paragraph_format.page_break_before = True

        is_heading = paragraph.style.name.startswith("Heading")
        if is_heading:
            paragraph.paragraph_format.keep_with_next = True
            heading_size = {
                "Heading 1": 16,
                "Heading 2": 14,
                "Heading 3": 12,
            }.get(paragraph.style.name, 12)
            for run in paragraph.runs:
                set_run_font(run, "黑体", "Arial", heading_size)
                run.bold = True
            continue
        if has_drawing:
            continue
        if has_display_math:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(4)
            continue

        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(24)
        paragraph.paragraph_format.space_after = Pt(0)
        numbered = has_numbering(paragraph)
        if in_claims or numbered:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if numbered:
                # Let the numbering definition provide the hanging indent.  A
                # direct zero first-line indent makes continuation lines start
                # underneath the number instead of underneath the claim text.
                properties = paragraph._p.get_or_add_pPr()
                if properties.ind is not None:
                    properties.remove(properties.ind)
            else:
                paragraph.paragraph_format.first_line_indent = Cm(0)
        elif text.startswith("申请人："):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
        elif text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(0.74)

        for run in paragraph.runs:
            set_run_font(run)


def restyle_tables(document: Document) -> None:
    for table in document.tables:
        table.style = document_style(document, "Table Grid")
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if row_index == 0
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    for run in paragraph.runs:
                        set_run_font(run, size=9.5)
                        if row_index == 0:
                            run.bold = True


def postprocess_docx(input_path: Path, output_path: Path, md_path: Path) -> None:
    document = Document(input_path)
    configure_word_compatibility(document)
    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.3)
        section.footer_distance = Cm(1.3)
        add_page_number(section.footer.paragraphs[0])

    title = md_path.stem.removesuffix("_CN")
    restyle_paragraphs(document, title)
    restyle_tables(document)
    size_images(document)
    add_figure_captions(document, markdown_figure_captions(md_path))

    properties = document.core_properties
    properties.title = title
    properties.subject = "中国发明专利申请初稿"
    properties.author = "【待填写】"
    properties.keywords = "发明专利; MSM; 无人机; 光电"
    document.save(output_path)


def omml_count(docx_path: Path) -> int:
    import zipfile

    with zipfile.ZipFile(docx_path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    return len(re.findall(r"<m:oMath(?:\s|>)", xml))


def build(md_path: Path, pandoc: Path) -> Path:
    if not md_path.exists():
        raise FileNotFoundError(md_path)
    output_path = md_path.with_suffix(".docx")
    with tempfile.TemporaryDirectory(prefix="msm-patent-docx-") as temp_dir_text:
        temp_dir = Path(temp_dir_text)
        reference_path = temp_dir / "reference.docx"
        pandoc_source = temp_dir / md_path.name
        raw_path = temp_dir / "pandoc.docx"
        processed_path = temp_dir / "processed.docx"
        create_reference_doc(pandoc, reference_path)
        pandoc_source.write_text(
            normalize_math_punctuation(md_path.read_text(encoding="utf-8")),
            encoding="utf-8",
        )
        subprocess.run(
            [
                str(pandoc),
                str(pandoc_source),
                "--from=gfm+tex_math_dollars",
                "--to=docx",
                f"--resource-path={md_path.parent}",
                f"--reference-doc={reference_path}",
                f"--output={raw_path}",
            ],
            check=True,
            cwd=ROOT,
        )
        postprocess_docx(raw_path, processed_path, md_path)
        shutil.copy2(processed_path, output_path)
    return output_path


def main() -> None:
    pandoc = find_pandoc()
    print(f"Pandoc: {pandoc}")
    for md_path in PATENTS:
        output = build(md_path, pandoc)
        print(f"{output.relative_to(ROOT)}（OMML公式 {omml_count(output)} 个）")


if __name__ == "__main__":
    main()
