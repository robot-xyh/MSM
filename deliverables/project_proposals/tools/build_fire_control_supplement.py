#!/usr/bin/env python3
"""Fill the fire-control sections of the proposal template without changing it."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[3]
PROPOSAL_DIR = ROOT / "deliverables" / "project_proposals"
ASSET_DIR = ROOT / "deliverables" / "图片素材"
TEMPLATE = PROPOSAL_DIR / "建议书模版.docx"
OUTPUT = PROPOSAL_DIR / "建议书模版_智能化火指控补充版.docx"

BODY_FONT = "仿宋_GB2312"
HEADING_FONT = "黑体"
LATIN_FONT = "Times New Roman"
BODY_SIZE = 14
CONTENT_WIDTH_CM = 15.0


METRICS = (
    (
        "目标规模",
        "面向200个来袭目标、200个拦截资源，按不少于8个区域实施分层调度；支持目标与资源数量不等时的任务分配及高威胁目标的多资源保障。",
    ),
    (
        "态势更新",
        "目标威胁和区域需求更新周期不超过1秒；信息不完整时，目标保持待确认状态。",
    ),
    (
        "滚动分配",
        "目标与资源数量均不超过50时，95%的分配计算在100毫秒内完成，单次最大耗时不超过250毫秒；200对200规模下，95%的分层滚动更新在2秒内完成。",
    ),
    (
        "计划约束",
        "已发布计划满足资源容量、机动可达、备用保留、空域限制、任务有效期和版本要求，硬约束合规率达到100%。",
    ),
    (
        "高威胁保障",
        "资源条件满足时，高威胁目标漏分配数为0；资源不足时，短缺原因及补充需求报告率达到100%。",
    ),
    (
        "计划稳定",
        "旧版本计划拒绝率达到100%，版本回退次数为0；稳定阶段5秒内非必要任务调整不超过2次。",
    ),
    (
        "智能调度",
        "独立测试场景下，综合任务代价较固定的确定性基准方案降低不少于5%，硬约束违规次数为0。",
    ),
    (
        "群目标配准",
        "单站航迹准确度和覆盖度均不低于90%时，20、40、60目标场景的双光电关系准确度不低于90%，目标覆盖度不低于80%，95%的关联计算在1秒内完成。",
    ),
    (
        "末端交接",
        "目标框宽不小于20像素、信息时效不超过100毫秒、标定重投影误差不超过1像素时，中心航迹与机载航迹配准准确度不低于95%，覆盖度不低于90%，错误全局锁定率不高于1%。",
    ),
)


def set_run_font(run, *, size: float = BODY_SIZE, bold: bool = False, heading: bool = False) -> None:
    run.font.name = LATIN_FONT
    run._element.get_or_add_rPr().rFonts.set(
        qn("w:eastAsia"), HEADING_FONT if heading else BODY_FONT
    )
    run.font.size = Pt(size)
    properties = run._element.get_or_add_rPr()
    complex_size = properties.find(qn("w:szCs"))
    if complex_size is None:
        complex_size = OxmlElement("w:szCs")
        properties.append(complex_size)
    complex_size.set(qn("w:val"), str(int(size * 2)))
    run.bold = bold


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def format_body(paragraph: Paragraph, *, first_indent: bool = True) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(28)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    paragraph.paragraph_format.first_line_indent = Pt(2 * BODY_SIZE) if first_indent else Pt(0)
    properties = paragraph._p.get_or_add_pPr()
    indent = properties.get_or_add_ind()
    indent.set(qn("w:firstLineChars"), "200" if first_indent else "0")
    indent.set(qn("w:leftChars"), "0")
    indent.set(qn("w:rightChars"), "0")
    # Inherited document-grid snapping otherwise stretches table and figure rows.
    grid_snap = properties.find(qn("w:snapToGrid"))
    if grid_snap is None:
        grid_snap = OxmlElement("w:snapToGrid")
        properties.insert_element_before(grid_snap, "w:spacing", "w:ind", "w:jc")
    grid_snap.set(qn("w:val"), "0")
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.widow_control = True


def format_heading(paragraph: Paragraph, *, level: int = 3) -> None:
    # The template's heading styles carry automatic numbering. Use its body
    # style here so explicit section numbers are not duplicated in Word.
    paragraph.style = "定义正文"
    format_body(paragraph, first_indent=False)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        set_run_font(run, bold=True, heading=True)


def write_paragraph(
    paragraph: Paragraph,
    text: str,
    *,
    style: str = "定义正文",
    heading_level: int | None = None,
    bold_prefix: str | None = None,
) -> Paragraph:
    clear_paragraph(paragraph)
    paragraph.style = style
    if heading_level is not None:
        run = paragraph.add_run(text)
        set_run_font(run, bold=True, heading=True)
        format_heading(paragraph, level=heading_level)
        return paragraph

    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        set_run_font(prefix, bold=True)
        suffix = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(suffix)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    format_body(paragraph)
    return paragraph


def insert_paragraph_after(
    doc: Document,
    cursor,
    text: str,
    *,
    heading_level: int | None = None,
    bold_prefix: str | None = None,
) -> Paragraph:
    element = OxmlElement("w:p")
    cursor.addnext(element)
    paragraph = Paragraph(element, doc._body)
    write_paragraph(
        paragraph,
        text,
        heading_level=heading_level,
        bold_prefix=bold_prefix,
    )
    return paragraph


def insert_figure_after(
    doc: Document,
    cursor,
    image_name: str,
    caption: str,
    *,
    width_cm: float = CONTENT_WIDTH_CM,
) -> Paragraph:
    image_path = ASSET_DIR / image_name
    if not image_path.is_file():
        raise FileNotFoundError(image_path)

    image_p = insert_paragraph_after(doc, cursor, "")
    clear_paragraph(image_p)
    image_p.style = "Normal"
    format_body(image_p, first_indent=False)
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.line_spacing = 1
    image_p.paragraph_format.space_before = Pt(6)
    image_p.paragraph_format.space_after = Pt(3)
    image_p.paragraph_format.keep_together = True
    image_p.paragraph_format.keep_with_next = True
    image_p.add_run().add_picture(str(image_path), width=Cm(width_cm))

    caption_p = insert_paragraph_after(doc, image_p._p, caption)
    caption_p.style = "Normal"
    format_body(caption_p, first_indent=False)
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.line_spacing = Pt(18)
    caption_p.paragraph_format.space_after = Pt(6)
    caption_p.paragraph_format.keep_with_next = False
    caption_p.paragraph_format.keep_together = True
    for run in caption_p.runs:
        set_run_font(run, size=10.5)
    return caption_p


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, width_cm: float) -> None:
    width_twips = int(width_cm * 567)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def format_cell(cell, text: str, *, header: bool = False, center: bool = False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    format_body(paragraph, first_indent=False)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center or header else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_before = Pt(1.5)
    paragraph.paragraph_format.space_after = Pt(1.5)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, bold=header, heading=header)
    if header:
        set_cell_shading(cell, "F2F2F2")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_fixed_table_layout(table) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def fill_metric_table(table) -> None:
    while len(table.rows) < len(METRICS) + 1:
        table.add_row()
    while len(table.rows) > len(METRICS) + 1:
        table._tbl.remove(table.rows[-1]._tr)

    set_fixed_table_layout(table)
    headers = ("指标项", "指标")
    for col, text in enumerate(headers):
        format_cell(table.rows[0].cells[col], text, header=True, center=True)
    repeat_table_header(table.rows[0])
    for column, width in zip(table.columns, (3.0, 12.0)):
        column.width = Cm(width)
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        set_cell_width(row.cells[0], 3.0)
        set_cell_width(row.cells[1], 12.0)

    for row_index, (name, value) in enumerate(METRICS, start=1):
        format_cell(table.rows[row_index].cells[0], name, center=True)
        format_cell(table.rows[row_index].cells[1], value)


def build(output_path: Path = OUTPUT) -> None:
    doc = Document(TEMPLATE)
    original = list(doc.paragraphs)
    if len(original) != 24 or len(doc.tables) != 1:
        raise RuntimeError("The proposal template structure has changed; refusing to overwrite a different layout.")
    if original[3].text.strip() != "补充修改":
        raise RuntimeError("The metric placeholder was not found.")
    if original[8].text.strip() != "补充，重点在这一部分，图文。":
        raise RuntimeError("The fire-control placeholder was not found.")

    # Retain the template's numbered headings; discard only drafting placeholders.
    for index in (0, 3, 13, 16, 17, 19):
        element = original[index]._p
        element.getparent().remove(element)

    write_paragraph(
        original[4],
        "火指控面向200个来袭目标和200个拦截资源，采用区域调度、目标分配与末端执行相结合的分层架构，开展威胁评估、资源配置和滚动重规划，主要技术指标如下。",
    )
    fill_metric_table(doc.tables[0])

    write_paragraph(
        original[6],
        "系统由侦察探测、智能化火指控、低成本拦截无人机和拦截效果评估等部分组成，面向走廊式防御与要地防护两类任务，构建探测、指控、拦截与评估相衔接的群反群技术体系。侦察探测形成目标航迹及威胁属性，火指控据此组织区域资源调度和具体目标分配。拦截无人机在外部粗引导下进入目标预测区域，通过机载光电完成搜索、目标核对和末端导引。执行结果及剩余资源状态返回火指控，供后续任务调整使用。",
    )

    write_paragraph(
        original[8],
        "多方向、多波次群目标条件下，火指控需要根据持续变化的目标态势，协调各区域的资源配置和具体目标分配。调度开展时，目标位置和身份信息往往尚未稳定，后续还需随航迹更新、资源消耗和任务进展调整计划。研究重点是兼顾响应时效与计划稳定性，减少重复出动、航迹交叉和指令冲突，使侦察信息能够连续用于资源调度、目标分配和末端目标核对。",
    )
    cursor = original[8]._p

    def add(text: str, *, heading: bool = False) -> None:
        nonlocal cursor
        paragraph = insert_paragraph_after(
            doc, cursor, text, heading_level=3 if heading else None
        )
        cursor = paragraph._p

    def figure(
        image_name: str, caption: str, *, width_cm: float = CONTENT_WIDTH_CM
    ) -> None:
        nonlocal cursor
        cursor = insert_figure_after(
            doc, cursor, image_name, caption, width_cm=width_cm
        )._p

    add(
        "拟采用区域调度与具体目标分配相结合的组织方式。区域层汇总各方向的目标数量、威胁程度和可用资源，结合相邻区域增援条件，提出资源配额、跨区调动及备用安排。区域内依据分配方案建立无人机与目标的对应关系；执行端根据机载观察持续核对任务目标，并反馈任务进展。"
    )
    add(
        "智能调度侧重处理区域之间的供需变化及连续波次条件下的资源使用问题。图神经网络描述相邻区域的联系，时间序列模型反映态势变化，强化学习形成区域级调度建议。具体无人机与目标的分配由规则和组合优化算法完成，任务计划经资源、可达性及安全条件检查后下发。区域调度与任务执行的关系如图1所示。"
    )
    figure(
        "ChatGPT Image Aug 5, 2026, 05_03_48 AM.png",
        "图1  区域资源智能调度与任务执行流程",
    )
    add(
        "目标身份管理贯穿侦察与执行过程。双光电通过航迹配准建立统一目标关系，交汇定位后形成可供调度使用的目标航迹。无人机进入目标附近时，将中心航迹与机载局部航迹关联；多架无人机同时观察相邻目标时，进一步核对不同视场中的对应关系。各设备保留本地观测编号，统一目标编号由中心维护。"
    )
    add(
        "任务执行中的目标变化、资源故障和安全冲突及时反馈至火指控，触发相关区域的计划复核与调整。一般航迹波动经连续观察后再决定是否换令，局部变化仅调整受影响的任务。计划注明生成时刻、版本及有效期，执行端拒绝旧版本、过期指令和资源重复占用的任务。"
    )

    original[11]._p.getparent().remove(original[11]._p)
    write_paragraph(original[12], "1. 系统方案与技术设计研究", heading_level=3)
    write_paragraph(original[14], "（3）智能化火指控研究", heading_level=3)
    write_paragraph(
        original[15],
        "面向多方向、多波次群目标，研究区域态势表征、资源需求估计与滚动调度方法，统筹当前任务和后续波次的资源配置；研究兼顾威胁程度、机动可达性与计划稳定性的目标分配方法；研究双光电、中心与机载光电及多机之间的航迹配准，保持目标身份一致。结合图神经网络与强化学习开展模型训练和独立场景测试，通过仿真、半实物及受控外场试验验证任务准确性、目标覆盖度、处理时效和安全性。",
    )

    original[18].paragraph_format.page_break_before = True
    write_paragraph(original[20], "1. 区域滚动的调动分配技术", heading_level=3)
    write_paragraph(
        original[21],
        "按来袭方向汇总区域目标态势与资源状态，估计当前及后续波次的需求。利用图神经网络和时间序列模型描述区域联系及态势变化，强化学习提出资源调动建议，匈牙利算法或最小费用流完成具体目标分配。计划统一检查资源、可达性和安全约束；目标变化或资源故障时局部重规划，模型输出异常时采用确定性方案。",
    )
    write_paragraph(original[22], "2. 群反群目标配准技术", heading_level=3)
    write_paragraph(
        original[23],
        "将双光电的单站航迹转换到统一时空基准，通过共面筛选、图神经网络评分和一一匹配建立关联，再进行交汇定位。中心航迹按拍摄时刻投影到机载图像，与局部航迹进行多帧核对；无人机之间依据空间视线与运动信息配准。统一目标编号由中心维护，证据不足的关系保持待确认。双光电处理流程如图2所示。",
    )
    cursor = original[23]._p
    figure(
        "ChatGPT Image Aug 17, 2026, 11_31_11 PM.png",
        "图2  双光电多目标轨迹配准与交汇定位流程",
        width_cm=14.0,
    )

    doc.core_properties.title = "项目建议书智能化火指控补充版"
    doc.core_properties.subject = "区域滚动调动分配与群反群目标配准"
    doc.core_properties.comments = ""
    doc.save(output_path)
    print(output_path)


if __name__ == "__main__":
    build()
